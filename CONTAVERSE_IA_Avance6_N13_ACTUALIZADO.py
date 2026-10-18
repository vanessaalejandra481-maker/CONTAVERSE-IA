import csv
import io
import json
import base64
import html
from collections import Counter
from datetime import datetime
from pathlib import Path

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

st.set_page_config(page_title="CONTAVERSE IA | Multiagente", page_icon="🏙️", layout="wide")
BASE = Path(__file__).resolve().parent
LOG = BASE / "decisiones.jsonl"

st.markdown("""
<style>
.stApp {background:linear-gradient(145deg,#090f25,#101b3b 62%,#0a1230);color:#edf3ff}
[data-testid="stHeader"] {background:transparent}
[data-testid="stSidebar"] {background:#0c1734;border-right:1px solid #304269}
[data-testid="stSidebar"] * {color:#e9efff!important}
.stApp h1,.stApp h2,.stApp h3,.stApp p,.stApp label {color:#edf3ff!important}
[data-testid="stTextArea"] textarea,[data-testid="stFileUploader"] {background:#17264a!important;color:#fff!important;border:1px solid #5369a1!important;border-radius:12px}
[data-testid="stTextArea"] textarea::placeholder {color:#bdc8e1!important}
[data-testid="stRadio"] label p,[data-testid="stRadio"] label span {color:#edf3ff!important}
[data-testid="stAlert"] {border-radius:12px}
[data-testid="stAlert"] {background:#1b3051!important;border:1px solid #6b83b7!important}
[data-testid="stAlert"] p,[data-testid="stAlert"] span {color:#f7faff!important}
[data-testid="stAlert"] svg {fill:#c9dcff!important}
[data-testid="stDataFrame"] {border:1px solid #3b5283;border-radius:12px}
.cv-kpi {background:linear-gradient(145deg,#17294e,#1c2550);padding:18px;border:1px solid #405a93;border-radius:15px;min-height:116px}
.cv-kpi .label {font-size:13px;color:#b9c9eb}
.cv-kpi .value {font-size:25px;font-weight:800;color:#fff;margin-top:9px;overflow-wrap:anywhere}
.cv-kpi .sub {font-size:11px;color:#a9bde6;margin-top:5px}
.cv-gabriela {display:flex;align-items:center;gap:18px;background:linear-gradient(120deg,#182b51,#292457);border:1px solid #6274b5;border-radius:18px;padding:20px}
.cv-avatar {width:155px;height:166px;flex-shrink:0;border-radius:16px;object-fit:cover;object-position:center top;border:1px solid #8b86cd;box-shadow:0 8px 24px #080e2c}
.cv-gabriela .name {font-size:22px;font-weight:800;color:#fff}
.cv-gabriela .desc {font-size:13px;color:#d3dffa;margin-top:6px}
.cv-gabriela .note {font-size:11px;color:#b6c7e7;margin-top:8px}
.cv-section {font-size:12px;letter-spacing:2px;color:#f4d281;font-weight:800;margin:14px 0 8px}
[data-testid="stVerticalBlockBorderWrapper"] {background:#142449;border:1px solid #304c81;border-radius:16px}
[data-testid="stVerticalBlockBorderWrapper"] p {color:#f2f6ff!important}
div.stButton>button,div.stDownloadButton>button {background:linear-gradient(110deg,#3757d7,#7c4de2)!important;color:#fff!important;border:1px solid #8899e9!important;border-radius:10px}
div.stButton>button p,div.stDownloadButton>button p {color:#fff!important}
[data-testid="stMetric"] {background:#15264b;border:1px solid #354d80;border-radius:12px;padding:14px}
[data-testid="stMetric"] * {color:#fff!important}
</style>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("## 🏙️ CONTAVERSE IA")
    st.caption("Ecosistema empresarial inteligente")
    st.markdown("---")
    st.markdown("🏠 **Inicio**")
    st.markdown("📂 **Análisis empresarial**")
    st.markdown("🧾 **Contabilidad y tributaria**")
    st.markdown("🛡️ **Riesgos y auditoría**")
    st.markdown("📊 **Reportes**")
    st.markdown("---")
    st.caption("Versión demostrativa · reglas locales")


def leer_csv(archivo):
    texto = archivo.getvalue().decode("utf-8-sig")
    lector = csv.DictReader(io.StringIO(texto))
    if not lector.fieldnames:
        raise ValueError("El archivo CSV no contiene encabezados.")
    lector.fieldnames = [nombre.strip() if nombre else "" for nombre in lector.fieldnames]
    return list(lector)


def numero(valor):
    texto = str(valor).strip().replace(" ", "")
    if not texto:
        raise ValueError("Valor numérico vacío")
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    valor_num = float(texto)
    if not (-float("inf") < valor_num < float("inf")):
        raise ValueError("Valor numérico no finito")
    return valor_num


def dinero(valor):
    return f"${valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def analizar(filas):
    resultados = []
    for i, fila in enumerate(filas, 1):
        try:
            base = numero(fila.get("base", ""))
            iva_pct = numero(fila.get("iva_porcentaje", "19") or "19")
            if base < 0 or not 0 <= iva_pct <= 100:
                raise ValueError("Valores fuera de rango")
        except (ValueError, TypeError) as exc:
            raise ValueError(f"Fila {i}: revisa base e iva_porcentaje; deben ser números válidos. {exc}") from exc
        iva = round(base * iva_pct / 100, 2)
        total = round(base + iva, 2)
        soporte = str(fila.get("soporte_factura") or "").strip().lower()
        retencion = str(fila.get("retencion_definida") or "").strip().lower()
        alertas = []
        if retencion not in ("si", "sí", "no"):
            alertas.append("No se indicó si corresponde aplicar retenciones.")
        if soporte not in ("si", "sí"):
            alertas.append("No se confirmó soporte de factura.")
        contable = {
            "agente": "Contable IA", "base": base, "iva": iva, "total": total,
            "conclusion": f"Compra por {dinero(base)}; IVA calculado {dinero(iva)}; total {dinero(total)}. Clasificación contable sujeta a soportes y naturaleza de la compra.",
            "estado": "Cálculo preliminar",
            "asiento_preliminar": {"debito_compra": base, "debito_iva": iva, "credito_proveedor": total},
        }
        tributario = {
            "agente": "Tributario IA", "iva_porcentaje": iva_pct,
            "conclusion": "IVA calculado. Verificar responsabilidad tributaria, procedencia del IVA descontable y retenciones antes de contabilizar definitivamente.",
            "estado": "Por validar" if retencion not in ("si", "sí", "no") else "Revisión requerida",
        }
        riesgo = {
            "agente": "Riesgos y Auditoría IA",
            "conclusion": "Solicitar factura y evidencia de la operación." if soporte not in ("si", "sí") else "Se declara soporte de factura; verificar su validez y trazabilidad.",
            "estado": "Alerta documental" if soporte not in ("si", "sí") else "Verificación pendiente",
        }
        resultados.append({
            "fila": i, "descripcion": fila.get("descripcion") or "Compra",
            "contable": contable, "tributario": tributario, "riesgo": riesgo,
            "comparacion": {
                "coincidencia": "Los tres análisis requieren validación antes del registro definitivo.",
                "discrepancias_y_pendientes": alertas,
                "dictamen": "APROBACIÓN CONDICIONADA: revisar documentación y retenciones." if alertas else "REVISIÓN HUMANA: confirmar tratamiento y soportes.",
            },
            "modo": "Reglas locales (respaldo sin API)",
        })
    return resultados


def analizar_con_failover(filas, simular_fallo=False):
    """Demostración de continuidad: ruta primaria simulada y motor local de respaldo.

    No representa una conexión real con proveedores de modelos ni con Frame.
    """
    eventos = []
    if simular_fallo:
        eventos.append({"etapa": "Motor primario (simulado)", "estado": "FALLÓ", "detalle": "Fallo inducido para demostrar recuperación"})
    else:
        eventos.append({"etapa": "Motor primario (simulado)", "estado": "NO CONFIGURADO", "detalle": "No existe integración de API; se usa el motor local"})
    try:
        resultados = analizar(filas)
    except Exception as exc:
        eventos.append({"etapa": "Motor local de respaldo", "estado": "FALLÓ", "detalle": f"{type(exc).__name__}: {exc}"})
        raise
    eventos.append({"etapa": "Motor local de respaldo", "estado": "OK", "detalle": f"{len(resultados)} operación(es) analizadas mediante reglas locales"})
    for caso in resultados:
        caso["failover"] = list(eventos)
        caso["modo"] = "Respaldo por reglas locales (simulación de failover)" if simular_fallo else "Reglas locales (sin API)"
    return resultados, eventos


def leer_registros():
    if not LOG.exists():
        return []
    registros = []
    for numero_linea, linea in enumerate(LOG.read_text(encoding="utf-8").splitlines(), 1):
        if not linea.strip():
            continue
        try:
            dato = json.loads(linea)
            caso = dato.get("caso", {}) or {}
            comparacion = caso.get("comparacion", {}) or {}
            contable = caso.get("contable", {}) or {}
            registros.append({
                "Fecha": dato.get("fecha", ""), "Archivo": dato.get("archivo", ""),
                "Caso empresarial": caso.get("descripcion", ""),
                "Valor base": contable.get("base"), "IVA": contable.get("iva"), "Total": contable.get("total"),
                "Contable IA": (caso.get("contable") or {}).get("conclusion", ""),
                "Tributario IA": (caso.get("tributario") or {}).get("conclusion", ""),
                "Riesgos y Auditoría IA": (caso.get("riesgo") or {}).get("conclusion", ""),
                "Dictamen": comparacion.get("dictamen", ""),
                "Decisión humana": dato.get("decision", ""), "Justificación": dato.get("justificacion", ""),
            })
        except (json.JSONDecodeError, AttributeError, TypeError):
            st.warning(f"Registro histórico inválido en la línea {numero_linea}; se omitió del informe.")
    return registros


def generar_excel(registros):
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen ejecutivo"
    analisis = libro.create_sheet("Análisis multiagente")
    validacion = libro.create_sheet("Validación humana")
    azul, celeste = "17365D", "DCEAF7"

    def titulo(hoja, texto, columnas):
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
        celda = hoja.cell(1, 1, texto)
        celda.fill = PatternFill("solid", fgColor=azul)
        celda.font = Font(name="Aptos", size=16, bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center")
        hoja.row_dimensions[1].height = 36

    def crear_tabla(hoja, encabezados, filas, fila_inicio=3):
        for indice, nombre in enumerate(encabezados, 1):
            celda = hoja.cell(fila_inicio, indice, nombre)
            celda.fill = PatternFill("solid", fgColor=azul)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.alignment = Alignment(wrap_text=True)
        for numero_fila, valores in enumerate(filas, fila_inicio + 1):
            for indice, valor in enumerate(valores, 1):
                celda = hoja.cell(numero_fila, indice, valor)
                celda.alignment = Alignment(vertical="top", wrap_text=True)
                if numero_fila % 2 == 0:
                    celda.fill = PatternFill("solid", fgColor=celeste)
        if filas:
            ultima_columna = get_column_letter(len(encabezados))
            tabla = Table(displayName="Tabla" + str(libro.worksheets.index(hoja) + 1), ref=f"A{fila_inicio}:{ultima_columna}{fila_inicio + len(filas)}")
            tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            hoja.add_table(tabla)
        hoja.freeze_panes = f"A{fila_inicio + 1}"
        for indice, encabezado in enumerate(encabezados, 1):
            hoja.column_dimensions[get_column_letter(indice)].width = 48 if encabezado in ("Caso empresarial", "Dictamen", "Justificación", "Resultado") else 24

    titulo(resumen, "CONTAVERSE IA | RESUMEN EJECUTIVO", 4)
    decisiones = Counter(str(r["Decisión humana"]) for r in registros)
    for fila in [["Indicador", "Resultado"], ["Decisiones registradas", len(registros)], ["Aceptar", decisiones.get("Aceptar", 0)], ["Modificar", decisiones.get("Modificar", 0)], ["Rechazar", decisiones.get("Rechazar", 0)]]:
        resumen.append(fila)
    crear_tabla(resumen, ["Fecha", "Caso empresarial", "Dictamen", "Decisión humana"], [[r["Fecha"], r["Caso empresarial"], r["Dictamen"], r["Decisión humana"]] for r in registros], fila_inicio=9)

    titulo(analisis, "CONTAVERSE IA | ANÁLISIS MULTIAGENTE", 6)
    filas_analisis = []
    for r in registros:
        for agente in ("Contable IA", "Tributario IA", "Riesgos y Auditoría IA"):
            filas_analisis.append([r["Fecha"], r["Caso empresarial"], agente, r[agente], r["Dictamen"], r["Total"]])
    crear_tabla(analisis, ["Fecha", "Caso empresarial", "Agente", "Resultado", "Dictamen", "Total"], filas_analisis)

    titulo(validacion, "CONTAVERSE IA | VALIDACIÓN HUMANA", 5)
    crear_tabla(validacion, ["Fecha", "Archivo", "Caso empresarial", "Decisión humana", "Justificación"], [[r["Fecha"], r["Archivo"], r["Caso empresarial"], r["Decisión humana"], r["Justificación"]] for r in registros])
    for hoja in libro.worksheets:
        hoja.sheet_view.showGridLines = False
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
    archivo_excel = io.BytesIO()
    libro.save(archivo_excel)
    return archivo_excel.getvalue()


st.markdown("""
<div style="background:linear-gradient(120deg,#122b58,#202454 65%,#432e80);padding:28px 32px;border-radius:20px;border:1px solid #6676b5;margin-bottom:22px">
<div style="color:#e6c76a;font-size:12px;font-weight:800;letter-spacing:2px">ECOSISTEMA EMPRESARIAL INTELIGENTE</div>
<h1 style="color:#fff!important;font-size:40px;margin:10px 0">🏙️ CONTAVERSE IA</h1>
<p style="color:#dbe7ff!important;font-size:17px">Centro ejecutivo de análisis contable, tributario y auditoría</p>
<div style="color:#e6c76a;font-size:14px">Contabilidad · Tributaria · Costos · Riesgos · Auditoría</div>
</div>
""", unsafe_allow_html=True)

st.markdown('<div class="cv-section">CENTRO DE CONTROL · VISTA EJECUTIVA</div>', unsafe_allow_html=True)
left, right = st.columns([3, 2], gap="large")
with left:
    st.markdown("### ✨ Panel ejecutivo")
    st.write("Gestiona casos empresariales, revisa los análisis especializados y valida las decisiones antes de su registro.")
    st.caption("Demostración con reglas programadas. La conexión con modelos de IA reales está pendiente.")
with right:
    st.markdown("""
    <div class="cv-gabriela">
      <img class="cv-avatar" src="data:image/jpeg;base64,/9j/4AAQSkZJRgABAQAAAQABAAD/2wBDAAUEBAQEAwUEBAQGBQUGCA0ICAcHCBALDAkNExAUExIQEhIUFx0ZFBYcFhISGiMaHB4fISEhFBkkJyQgJh0gISD/2wBDAQUGBggHCA8ICA8gFRIVICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICD/wAARCAEzAYYDASIAAhEBAxEB/8QAHAAAAgIDAQEAAAAAAAAAAAAABgcEBQACAwgB/8QASBAAAgEDAgQDBAYIAwUIAwEAAQIDAAQRBSEGEjFBE1FhByJxgRQyQpGhwRUjM1JicrHRFiThQ4Ki8PEmNDVTY5Kywhdzgzb/xAAbAQACAwEBAQAAAAAAAAAAAAADBAECBQAGB//EADMRAAICAQQABAQFAwQDAAAAAAECAAMRBBIhMRMiQVEFMmFxFIGRofBSsdEVI0LxM0PB/9oADAMBAAIRAxEAPwBO8PHLKKYkKg260u+GxllpkQL+oWvU19TPJzObCtMV3Zd60xRJWc601DH0Q/Cu5XauGog/RDkdqq3UssUWs7a6nxpg8OY8JPhS+1n/AMcX40fcOFjEnwpar5jLNChwMVFbrXduY9q58p8qZlDI5FYBUpIGb7Oa2Nq46rVSRJCnuRPlXw1NSzkc7Ka6fo2TuKrvUGWFbHoSsIriyZPSrc2Dr1BrQ2bDqKuHEqam9pUiI56VuqYNWsVjLPMsEERkkboo/r6D1qn1rX9H0RXghk+nXy9TGOZUPpnb5n5CltRrq9Nw3JPoO4anRvd1wPedZ0xb8zkIjdGc8oPzNUUo0dpOW81mC2Q7cyxtKfuGP60GapxFq1/M0ioFB+08hY/M1Rs01w5N0eYn7SHOPlWU/wARvt+XCj9TGxo6azzlj+gjYh4b0m+iX9Ca6l/M2TyeCIwBtnq2SaH7yM6NeNFqCmB16gruRtv8N6D7ON4pC1reGNhseUkY/tVpq2papqkj/Tm8SaZVj5h2A5ebHxCrmgLqtQh5OfyhWoocdYMJbPifRI3AluWX18P/AFox0rVNL1BR9Ev4ZCfsluU/ca8/XkEwmMb7DOSR0/6VEF3NFIv0YFFX6vmfWm1+IWDsCKNpk9J6l8IqcEEfGsKCkRoPtE1rSXSN28eEfWikOQR8O3yxTf4e4p0riVAtpJ4N3jLW0je9/un7Q/GnatXXZx0YE6dl5HMtSu9ZnHWpgtWbtXw2TnsaZ3iDNRmmnzY1CMDzo8kkLWB3+zQVZWbrfRnl70aSqVsj/LS9zA9Q1aFe4ob/AH4rH81PPhNf8iv8tIm+f/tYB/FT54S3sE/lpe08Q1HZg5x6P8lL8KSop2cfj/JyY8qSoFVX5ZdvmmY901QagnMxoh6KaHdTkCk71x6lZHt4VxmrCIKCOlU0d1yr1rqLzG+aSfmHDgQjjdAu5qRHIM7HNCh1PA61Kg1LK9aQOfWFS7mevPZZdxx8OQAkfVFX/G2tw2ujyMXA286898E8a3FjaRWaks56AVccV6jrOu2JiAZIz1ry9qOHKHombqFGxYvcEtZ4vSS7ZVfIzQ7eaml0hy3WpzcMMsZZwc+Zqgv9LeEkJkU7VsXhYtYbezKDUo42YlcE1VyKyjYVZTwXCybjIrGtuePPenAwMzGTcYNTOS+K7QjA2NbXVqyyFgD1rtaWkkpGxopIxAhTmc2kYHArKvIdIyuSuayg+IIyKHMuuFoWdlwKaNvZyG3XAoM4Msw3JtTls7BPoq5HavZnUheInVpNwyYHNYyZxy1Lh0NpF3zmieSxRGzipVuY4xkig26xgPLGE0ag+aBh0aRWIIOxqv1uxaK0O3amEUSRy2BvVHxJBGLRj6VA1bHGZZtGgBInmnWQf8QIvrTX4Q03xYUOO1LHXFX/ABMuP3qdXBAXwI9uwqXsZASIpp61ZsNLaXSVRdxVe1iDJgCi68HunAqpSEmbJG1TXcxGTGLaEBwJ107RlaPPLXS50hRn3aurFlSLavlxKpOKQNtniR1aq9mJUW2lKOq1Ik0+NewqdHMgWo01yCSBUbrGaWKoq8Stlsl7LXCHSpby5W3gQcx3LNsqKOrE9gKtIEmvLqO2t0LyynlUVC43v7TR9Jk0C0mZ3nX/ADMkbcrTHuvN9lB5jrQ9VrvwqYHzHr/MpTpxe+PQdwC4w4viitp+H+EVDxgct5qLnkD+Y5uy+Sjc0phaSTFnRzOoOWmccsefQd/nRpPpUR09dU1dvB01M+Bbp7vjH0Hl5nvS+1riGTULn6PYRKI12VV2SMfmaydOzWsWJyfUxjVba1A6HoJtdXFlajnllM7jv0UfCqWfWLd5ABbAr5hsGuU8UUOJLktPO3QE/lUdraY+/cBYh2U7H7q0eBMkkmdmuGglS5tpi8bdCevqpHnW8uszsI4kY8wcFTnoPKoQjIQqqu4PXBH9K6paeLZSTocvER9xqQZXmWsitNbs5mUZPKXJ6nvgVUSWzoW5WIz1Pet7QyROOWBpWHVuYDHwzVmZCQPFiMZPTmXGfmKuBmVg+EKnAYfI1Psbma3lWRC6MpyHQ9D57dK+3UCscqvK3mKiIxR/eJUjuKoRLKcGP/gT2m2t00Wl8USqjHCx6h2PkJfL+b76daaXEVDe6QRkEHII88+VeKreRSvM65P/AJkZ3+Y6Gmz7OvahdaCIdE1uUz6Qx5YbjOTbHy/k9O3ao8ZxxmOV7D2I/V02NJlIArfU18O1bHlUSHVElmTlcMrbgg5BFSNTlDWrZPanaQ/bQWpKdCJS7bPGAH8VegeEd7FP5a893BzxfkfvV6G4QH+QT+WnLvliFHZg/wAfD/JyfA0ku5+NO/j8Ys5PgaSI6n41RflhH+afH+ofhQlq7crE0XPjkO3ahDWfrNXMOIJpSLKSa3Mh6io2DXaNWYEAEmlGXiRiRnnYPg1Jt5223rlNazEgiM/dUu0tJCRzRmknGJdATHX7HtGGqXTzvEWC7cxFPy94bgTTTiMDbyoV9iWkwwcNQyBMM4yfWmxq8YXT2wK8hqSXdm9BPUUEVBU94h9fsVhRlRQKX1zZF3YsO9NniJQWbbNAU8QMp2q9DcR2xcwJvtOjVckUO3UQifC9KOdUTCHIoTuoS+SBW7Rp2YZmJqGUHAgxPymbDDarbS44WIziqjUkaMlh1FQbbVXtzvmiWUHGBF6nAbJjKht4uX3ayqTS9TM0PNntWVmshBxNdXUjMJeDboApvThtbwC2TftSJ4PkPMnxpv2zMbZa94aQRPN1aggYly90GYAnrXZJlPu1RsW6+VYt06tigPR7RhdRnuW/jAEgHvVBxNOTZtv2qWspIzVLxHIfoh+FFFIzBNqTiIfWWzxGhP71OvgmUCCPPkKSGsb8QofWnDwa5EUfwFTs3ZEUrs2nMY0zqwxUPYNWxbNfMd64VgQzXEyfat+qzUW8ciXY9q6QPyRYqLcvzyfKhirzSTccTkbggda5rI8sqogyx9cfM+Qrm6nsM5q84b0GTVdVELrm0iINw3aQ/wDl5/dz186i5kpQu3pLIz2MFEtLd4eHeFm1UxmW8vv1VsvQyZ8vJe5PXA9aW8OknVby81TV5TPZwNi4cHHjyDpAnko25vLpTD4k8bXOI47KxkMYjiISTG1vDnDS4/ebfHwFKX2u8RQ6Jo0HCOi/5eSSPDAHeGLux/ibua+fPZZqtRn1b9hPVoqaenJ9IpuPeLLniPW5LO0cC0iPh/q9lAHRV9BVJbWYgRERMs26p5+p9K30yxVVadkJjXov7x7D+9H2n6JHo+jHiLW4wPE/Yxtt4h7Y8lFeh8mnrCL/ANzCAfU2F2/6gomiR2cBv9Qk5ZX3BPX5eXxocv7uCNmWABM/aKFvxrtr3Ec+pXzyJJhM4DYz9w8qo0aWZ9pHfP7woiKx5aCsdflTqc2eRpAc58uWrqxWT9G3jtGxYoAT6Z6/Gpek8PzXJ51j517jG3+lGMHC/wBKsXhIaOQ4yp6nH9RVmcCVSpm6i3tIInk5pyqL/GaIYbdXh5ba8jmQ/wCyfp8s1NvuBrmB+ZST5BqoJbTU9MlyhdceW/8AWuWwHoyTUy9idLmzdCeReT+BvyNU08TByCMHyO1Fmn6hDfxeBeoiuNucDAz5MO3xqDqdgsLcjgmPoG6lPj5iibveCKeog5DJJC+VOR3Har6xljm2GzHqPM/3/rVRLAySFW91x0PY1tbsQ2Pqt3B6GqMMyyEg4jp9nvFzWl9BoWpzYt88tvK/+yP7jfw+XlTt1Hm+iPnyryPZXbeLGSxEkZBVzuR6HzFekuFeIYdY4ShRj+sQGLBOSjAfUPmMbqfLbtTulv8A/W0rcmRuEAX/AP8AWjP71ejOEB/kE/lrzvMB/jED+KvRPCO1gg9Kdt+WK09mD3tAH+Tk+BpHn6x+NPL2gD/JyfA0jftH41VPlhG7mN9Q/ChTV1yTRW31T8KFNW2JzUnqUMoxHkgedHPDvD63EankyTQSrAMp9abHBVyvhoDVUGZVpzu+GFjH7MD5VzttCQbcgo41RlbceVVFqR44+NJ6weXMd0oBOI7fZTEItAijAxy7Ue61/wBwag72ZrnSAfU0Z62h+gMa8TYPK82j/wCZR9omeIR7zUFsuZWyKOOIFwzUDzMFlNCoPE03lJrKoITmhO4K8hwaINenIhbFBEt6CDk17LSsPDE81qhiwyi1Zt2occjNXeqShs0PjLTqo7mufkxYQz0JSLb5VlTNDiAtt/KsrGsbzGbNanaJccHLlk+NOC1GLZaUXBv1k+NOG1XNquPKveE8TzCDifGGx2qPyEsTUzCc4VmAzW/PaA42oLuFIEcrr3AmR0U8tU3ES/5Q/Cr1JIySqkbGqPiTP0VvhRgYow4iG1hf+0CfzU3uEBiKPHkKT+st/wBoU+NOLgz3oo/gKGp5MGozDgZrdQTW/IK4T3MdsuSd6sDmEK4k+K35o+YnAqPPByMN8g1RzcUwRIULgEGo8fEy3c0UMR52ZsADuaXBffz1HTXWagR3CW2g55dhuoyPQ9BTC5bbhrg4+MfCZ0LykdQgGW+eNh6mhTg2BNV1VQn6y2gbnlk7Ow7D0HT7zU3iK/8A0prYhY/5OFvGYHYGOM+6D/NJj/215f43qskUr9zNP4bp+S5kd7g6Fwvd6zqCql1c5uJgekeB7kf8qKPvrx/q2p3fE+vXeqXBZ5b2UsAT0TPup/T8ae3ty4na24YtdAjm5bi+3kHcRA75/mOaTHD2ll5EYRNI0kgVIx1kfso9B1JpT4bUERrm9evtC65yziofnDf2f8Iw6vqwmvABpenL407HpIR/c7UM+2HiY6rxI+nxHks7JRGUTYcx35R8NgfhTWmu7bhXgO3slYGe6Q3twRtzRrsg+BI29K82ag0+o64TKSWc+I582Y5J/GjUk22m1uhwIO8CmoVL2e5GsdMutRkAjj5VPYUc6XwiIoh4sLyk9lH96u+EtEDLGQm/wpu6RwrbuBJNEHY/vb/hRHvOcCWp0igZaLjhvhG4aQvyPHETvhtqYttoEEQXlhAx6UbW+jRxQhVjAA8hX17AIdhtSzszdx6tEXgQHveH7a5Q+JGc4wCO3w8qCte4UVoDiNWIGc46n/WnFNbKFIxvVBqFnzqRihhiDCsgYTzTqeiPbSmaAlXXuOvwP9q0jb6VaLHMAG+r8D5fCm7rWhRy8zGEZPcClpq+lGxMrxqQp6rWhXbuGDMm6jadwgtPaFWMDrn93P8A8f7VXPGytynPMOhq7kmF5AWJ/Wp18yOzfGuSQx38TZwsy/WHn/EPQ96OGx3EmTPUjWchdlVvddeh7f8ATzpg8G6ydL1NUldktZ2VJBn9m32W+Xn5E0u2UwTeHN+rkG4PQN/Y1c2F14ihW91x7oPY+h8vSqsSDuEJWQfK0bOoW7RcWJOv1XYAgdj1+6n7wixNinwrzlw5qB1QW6XJ/XwAQuT1wPqn5dK9HcIjFigPUCtlbPErDRI1+G5Eo+PwTZyfCkbjBO3enpx7/wBzk+FItvrH40VPllG7mEjlNCWsNuaJ5WwhxQfqzkuc1JECxxKsNtTC4SuQqKS2KAoI1Zd6vNNvhaEDOAKlEI5gy4Mat5eKYlPP2qDb30ay7tQhNryyIFVs1xjvZ3OUUn4UvqU3jBjent2nM9bey69t/wBCIOcZNG2u3tsmmuzOPvrypwPxXrWmMYo7Z5Ys7AUUcR8XcS6hbeFHYvCh6ljXir6HR2ryMGehTbZi3niXXEOq2rFsMDS7vtXhWUkEVT3Da1cORI2M1EOiXs/15TRqtNUg8zS732N8qzXVtRS4RsY3oLlXMjY86OV4Tu5dh4j58lNT7T2fahLumnXEhPlGa069Xp6lxmIWae605xFRcWpkztUWPTiJQ3Ifup7w+y3WZfq6NP8ANMf1qxg9j2vSYxpip/M4FQ3xLT+kH+As9TE1ZPJBHy8jVlP629imsOmZEtk9C2fyrKTOtoJ+WNChwMbxFvwtw1q9oyma2IAo7uL76Fa8rjlKjfNGpFlHGvh8uw7UueMp4yWEXX0r0vw/WtqvnER1ekXTrlTKDUOKvCkJD9Kqm4yYtsTmha5jnlnc7kZqOLaUfZNbbVKe5iLqXXgRk6VxEJSCz7mrnV7lLiwznOVpTWzzW7gjIorOoO9goY/ZqGHrLI24YMXWtr/2iT+anFwYMQx/AUnNWfn1+I+tObgsZhj+FLA8mErGTDt35Y2NA/EOoyRq/KaN50PgMcdqBdTs2uZimO9SrgcmFsrYjAi+urm5mZmyQKnWKXsDw2cPMb+9AG3WGNunwZvwX40RHQ44IpLuaLnht18Rlx9c9l+ZrppljPbiTUbhea/uWznH1OY4AHqegHYCq2ahWXiDrodW5jv4ZS20bggrbygAkxGYDqQPfYfDoKp7NknWW7uPcikk8eTPRYYvqr82/oa01u4/RWi6focQy0EI51H2m8vmxx8qEuOtQl0D2aagI3IuLnFopH3ED589fNbnbUXls/McCe4qUU0jPoMxFcWa5Nxjxzf6iJCYFfCMeiRqcZ/P50aaFYJHqWnaNEOW8uykUnnbQtuR6OwyT5DAoI0hI9F0yO8njR5ZG54EYZE0g6OR3jTsPtN6Ci3g+eSyl1PiG8ZmawtnlDOcl5ZNlyfMkk1v2DbXtTodTEq81mW7M78d60NV4m1pYPdtYgLWEDoEQcoH9T8qXNpaCXXpSBsSMfcKm/SXmt5pWYkyyHJ89j/rXXRAJb1JAM5Aq4UV17RIz4lu6OPg7TlWKMlabOnwoqAAUA8Iov0dM0x7JCQMCkc8zWx5ZPRRy7da4zRc3apqxEDcVsEXO9WMEDiD89uc9KqLmDGfdownSPFC2v6xo+i2xl1C8jiz0TOWb4ChlSeoYOMcwavbIODkUAcRaOsiOAoxXLX/AGm3NzcG00CzKk7B2XmY/AVVW+m+0u9ia7MkiRvvyXBGD/ukUdamUZJxFmuVjtUExcajp8ml3/MoPLnb+1RZF8GRbm3bA6j0Hcf3ow1qw1cN4OsacIi3SaI5TPr5UFyiW2ma3l2390+Z/vTSEmZ9i7TxJcj2+o2xDjBHUD60Z9PSq5PpGm3C+KOeI7Bh0/59Kjs0sEoliPKRsD2/lNWVvfQzp4NzHy5GCCMj/p/SiYx9oDIY+xhhw7fxR3kVyrHBHK6/vL5j4V6o4HvFmsvDLczKiuD++hGzflXi+3iudPnEtk3iwk5MTHceoNekfZJxRDqNhDEcrcWTeG6tsTE+34Nj76Pp22krngy1nmXJHIhjx1lrOT4UjmU8x+NPLjVSbOT4Gk0YdztWqp4iDDJlfLETEaDNaXlLGmDLH+pbbtQNrUYLMMVO6CZILC8aPatheyHpmu/0MM2SKm22nqWHuilja3vLCoTtw7aXWsa5a6fH7vjPgsew716u4Y9kXDMmmqbmKeWXG7+JjJ+FeeOHSmk6rb3yICYmzjzFelOH/alw1BYKLi4eFsbqYyf6V5n4pZqGYeHnH0m3o0rVDnuXulezPR7G4b6EkgH8TZoguOC9PeILcQh1+6rThLWbLXLH6bZEvHJuCRjNX9/7lsWK5xWYtAeo2sTkQz6h0s8MdQOteAOGVPN+i4WPmwzVvDwjocOPD0y3XH/pioM/Ekdhnmt5Gx5VB/8AyDAW5Ftyp/iNBXUaUDzdwjVapuv7wqj0WwiHuW0a/BQK7/QLcDZFFB/+Ni5wvhjPrXVOJ55TtKg+Arjr9KvS/tK/g9Qez+8LRawjoorYRRKPqihJtbnbrcn5VGk1hj1uHPzqD8TpHyrJGhtPZhv+qHcVlAJ1ZQd5GPzrKH/qo/olv9Pb3nmbQuItQuo8PMT86n6hFJPAzuck0F8KTHk3o78GW4ttjyr519W01aV9DE81dY9i4zmDo0tOXJAzWRaJJOf1cWR51aXVrNbIGGWXzxVhpkuYkXmxinLrdqblidNJL4fiC8+ieHkMuCKjXNsYrbbsKK9WljaQ8hzjqRQ7qT4tNvKqB9yZMO9Sr8sWmqcq61ET2NOHgi6TEYztiktrTE6mKP8AgS8k8REz0pBj3LUHDAR9OFa1LelDHgwm8dpDhBRHAGfTR6ihy7sZmmDKSOY8vy7mkWsDArmbJrK4aWl5YJJp9vDDCpjZvEbm7ntn071x4Y0+HUOIRdynntNPzMMj9rL0DH8vICvjzXF7b3HvskCQlV5R9Ren3mrnRtKmseFxFCvLe6gxYA/7NBsCfIAfiax9TYadOyg8nj9e/wBozUgtuBPX+JBDtrHFFzdxKZI4JFSPbIaX7P3bsfhS29r+s6R9Ps9P8ZruKxUiK0XZZ36GSVuvJ1wBuxJ3xR5xhrtpwbw6dIsHRLkxkzSscFAdzk9mb7wK8x395Nq+pzahOTJJIRyKVxsOhI7DyFY+goLvvPQjmuvCJtHZnb6XPqGo/Srt/EkOO2AAB7oA6ADsBRDq18dO4IgsFOJtQk+lSefIPdT/AOxqg02ya5uI4y/KjHLyHsv2mP8ASvuu3h1HWFVV5YyVRV/cjA90fcB99bjgEhR6THrJVS59ZoB4dnapjGTkj5GrThSEvKmOmdqpLu4/zESnoqs5+6uWk6xqniRWujwu0vdkXJz5CudSwwJatwhyZ6l4XthDaxs+2RR7aXMMagcwzXlyK09o93brvep5frcYqXZaVxtFcCa4uLkP+8029KmkL/yEfF5bjaZ6mF0sg91s18kkwuc0peFtQ1u1uFW+k5kYYbLZ386ZEMz3EJ5d9qWJwcRkLkZg9xbr09haYgyHfIDeW1JSbRtX4ivDPeSsqMcl5Dkmm3xGuATIvNg7A0F69dyafpPNHzxvJsDGMuO2FH7xOw+Oe1Xrc+nch6wfm6mmnjhvhIHwoPpV+B73IvO4+PZfnVRe+1a1k3NvIkRYoDsdx8KAeJ9M1+PUpNI4pt77T7g+E9nZW7ARIjHLM46uxGBnOc5zXKDg7S24QVnt7j9MyTMYXRzgxAj669BvkDvTf4cDl2iI1bE7aljAk1O31lQEYNz0H8Q8NcqtmPGenpRN7P8AghrHF5dxu8o6czEgfKi3iXTFntsqgBA8qX3bWwpjmw2J5xzPNni/R5nt7td1OOfHb1rGtwQGhbK9QOv3Vb8WaSYNREi5UuO3nVIkEsVszI2GAyB2NPg5GRMVlKsVIk60u57ZgrL4i9eUb4/tTA4J4itbHiK1nFytpznwpGlGwRtve8wDg/KlVHq6FgJo2Rh9patLLUFkbCSLL6HqPv3ripBzOSwdZnsTiO8F5orOy8kqe5Kmc8rAefcHqD3BpVZyTtWnAusWt3wzNYvqc301G/7rI+UeIdOQncMN/jXeeEx3DKAQvbPlWnVaGUCBevacyPN+xb4UDax+0NH0q/qT8KBtZQc5opPEXc4lIjCpcEgBFQuUgbVgkZTvSvBkq8IYbnGN6mpesNuahmK4OcVKE7DfNJagbRmGFpE9hexjVon4ZghDjKjBpsX90htDk9q8d+ybiybTtX+gu5Ecm4+NejJ9ejksd5BuPOvKWag0bq/ebFdQ1G20SDrlwjK2MZpVa9fSW/M6EjFGepXyS5w+aAeISHgcis6irJyZs2cJgQSk43uIJSPEOxot4e4ya7ClpT99JnVsR3rA9zUzQr9oJgqtWm2lRl6mHVrLFt2seJ6NGt5j5ueoM3ESKxy/40spuIXitMc++KFbjiqWScqHOKVXQAzSt1ypiO8cRI3+0/GspJx6/cMPcJNZRfwIlRrwRM4SiBgzTLsPDMKI5xvS54TYC1zRhLcOluDGcEGvqKZfieVGEQGG8mm291AI1QEYoXl0Twb8xpIQp7A1PsOJOW3EUq4bHWo898skhkDHOe9Lbbq2I9I1uqtUGSJdKthaFREDtuTQVrlusSMqjaiuXWJfojIgBJ7mhPV5ua3YscmppNgJ3ylwrI8kUmuqBqYxRtwCM3SfEUEa4c6mKOvZ6hNynxFdYcAxPTj/AHBPQ9lHnT0GO1R5mjBfnUcqDkHr51Z6dAzWkYHUjaqGaR5NZ+jRqWRTyqMZLHz+debclmInqxhQMyy06N7xUhjtlSPmBIx2HnUzi7iS04U0x5CVa+dMopGfDHZm/Jany3kHDGiPLcRJ9JQAsHOFi8i58/4RvXmj2icYSatfTFJWZixyx6/M9vQCspw11gQQrOqrugdxNxDdazqlxPPIzhmJPO2QPMn1NUaO74jjRyXIGMe85PQenwrdLC6mVZ5MRRZ2eT3VHwHUmp5ltNLTnDnxyMBjs2/YDt/WtutVrUKsw7C1jFmneWdNJsPALBpNmnZd8kdEHoO/mfhVRp5Z5pLmX6xyfh/0FfG8W/kUkYUH3V7D1rrdR+FC1tGccozKw7Dso9TV8Y+8pnPXQkViLnx5c4Rh4an0/wCc0bcLTWWh2qTkIuBkk0FxxkeHEdiTzECpOpWesS2aiziLRDqwPSguN3lzGaTs82Mxi3ftaitnAiURxjbZOeR/gOgHqajye0HV/wBI2sNxp+pwtdJHLErxI/Oj/UYAdmwcb0OcHaZaPoWo6ZqWkXclxeEFbuEAmMruux7Z6+dHfDXBiQXNlqOpX8huIHWVIY12V1xy9fTO3QZqpSlBzDLZqHORxCrQbqXWInMERE8P7WMAqyHbqp3FMzhO+XPg3AyQMb1V6FoFvDdvrklpi+dSDcyMWlcEDYnyrpbq1vqfOu3M29IsAD5ZoqSww014sSM3QCgcuahXWgaRr1vp8k/i81nlhDGeUPJ2YnrsOnxqy4ltnKpKRnO9QNKcpIu+1QrESxQGRbnh+9uJlkNssjqvIjytzuq+QY7iuFrwY8bDxVihj68kYpiQIsiA1s9qQM9auSTIUQTFhBZw+HGuKo9RjWRWUiizUIyC2BihO/JDGhDuMlPLmJTj6yCyQ4G5lAHzpeKHjvbiADmjQ7n9002uP4maSydR1mU5oOvLG2h0x/ARY1QmWTJyzeZNPo2BiYtyZcmLFj+udXTKhiNuo+FTILXmJ3ztzIw2yK+ww+Mxl5fdZzUq0LQyFCMhWzT5PExlHMYeg2bWXCFtxPaMXSKYwXsB3PXaRT/UUwG1KCWwDcqyF1UpL3Ud8fHahLg6WJ+CdT05l5l+kBh5crAA1mjmX9GPasSfAcoD6A4qUGCpHrGGJwRLm51JFQgUH6jcrLIauLq3dkJGaE9SWSFiTmmn3EYmYzE9yVFGsi1pPakbjeq+01DkYAmryOVJ4+2azGd6m5nCVKArJg1OA90GtJYP1mRXeOJmTpXXOLFyIQGWug3TRalC0JxIpyDTNbiu9WAIzk/A0C8O6BLNchip33JHYUeTaTbW9uAFGw715+8IX5m5pEsVOJL0/WZrpSZM4ra/dJrdgTvVB9PgtImVGGarJuIF3XnoYq5yJoeOFGGMH+IdPZrjnXzqrtlFs+T1q8utQjuM5IND17KOYlTTwBxgzFuKh9wnXUNU/V8gPaqi2V7mYBe5rQqZH941ZWRjtyDRlTAgBm1smFWk6Sogyw3xWVWf4jFuAiNWUuanJzNZbqVGJ04XlIh5ewNGpk54BigLhrdcds0fRxFoV5RtX0KoANPOEkpJun6eblwWblHpVtNpCrAxWQnFQ7K5Foh8RSK7JrUchMJz7x6113iZyOoWgoBg9youdOu0iZ42PJQ3qqyRwYfJJFNVzbHTsAjpQFxHHCYjy4OKSqvZzhhGba1VcrEtref0kNu1MP2fDlvYB2Y0Ba8ANSGPKjvgAMCjE7seVfh3P5VW84UxbS82CeoNICvaIF7rUB9Lmg1hmtLkWbnd70rk2qHuo7yt0X90ZNb2l7FpmgS38/7O3iMhHngbD78Up77i3Vrnie7j1fVpoY4W/VWsT8kUfMASxA3Y743NeQsdixAnqymBzC/2g33DFrYW1qb1mNtkrHzFzzHqxH2m9TXm/WuINIguma1tE5ubZ5T4jfdWa1qGq8Uar4dvO6xSsxUZ5QkY7nHoMk0I3mmSwHnBPKT7oxuR2++j6fThfmPMz79Qx4USTd67c3kgaMHmxgO+5HwHauNnp891NzuWdz1YnpWWlhNzKJSFJ6/wj19avlvrSxhCRqp5R1Y4UfHzrSUBepnnLHLTdbY28PJByhlHvSNsqDzqma6gZyIXJhjy3iMPrHu5/Ku1xdXGqREvL4Vkp9+RhyqfRR3qkvryGQpb2ylIAe/Vj5n+1R2ZJYKJO0a7F7q0jFSqDCoD1A/vTp4d0qG5tVR1BDCkZpsZstRMn2RjPw86fnCdyPCiOcggUpqOORNHQ88N3DHT+FLZQFReUelFOm6DbWrZjhHP5nc180ibm5dtjRhbW0fhhzSY5mu6BJCKGO3K43xQ+7r9ORQd80TXkyKOVaC7uRLO6lvJ5QgjbcscADzqrHmUCwsvrQ3WnAlckCgmRbm0uysa82N8UVRcaWc2jLBE8LqPtIQc/Ol5dcT6hJxE0VtpIks12a4aTBJ8lUbn+lSTnqcAR3GFpGrRy2qsRgjYg9qt3v4jHt1oL0OOcxTTXEfheI3ME8hW93dz2rEqpkQfu9RVdxEuAplveuHBOdzQjqmADipS6t9JUlCW7fCq+dnllAbauXuXdsCL/j0La8OtqM0bMluwchep3x+dLLXeMNIuOD5tLsFdri6ljm5ygBjAUqyseuDsQB3yacHtLiU+z3VAQNoc/wDEK8suc4A6mtPToGGT6TzuttZGwPWXemBfo5ye+cfKusnu3/uj3WAOKjWw5LRmzjIIFWkMPizRFtzy023UQUZIjB4KUfoi8UjAMkZPwAyfwFXGiWjPDJJy/tVeY57AnI/CoHC6LNoz2kR/WzufEI+xGAOY/PoPjRTBEEilZfdBTlAHYHG34VFZLDj0jDDaefWVkqpykEUG67bhw3KKN7iPlU70JakwMhFO2syjMyzAZ7eWN84q1sJ2GAanSWqyLkCoLQtC+QNqzntFgwZEuVKyYOasoFiUDOKGkuio61JXUWCgZpIDHEsDHtw29nHZRSryklRQ7xvxELNmSDYsO1Bml8TSW8IjWQ4HbNRNXvDqbF2JO1Z66fFmWm2dVmnC9ytXWrieTlDEk1LKu0fOWJNQNN06RpSwG2aJ2sCluMjfFNPgHiL0ozjLQdMrxsQTUeWQsSTVnPZtznI3qC9qzHArgRAvW2ZXvLyHavoa4l2QHFTYdOMsw22os03RoygBQVLOFEvVp3cwPg0uaUFnyTWUyU0mOP7IrKCbo8NDxzBXh48gIPnTM0yMTrEqb5OKVekSHIHYmmxwvKi3Vvz9Aa98uQuZiJhsLDWPhdriAYhOMbmqCXhpre/MZHu+flTfsLy1FluwG1A+valbxXMjIQTWXXqbnYrNI01jmUt5pq29gVWQ9KXmtHkiIJyaO77W4nsioB5sUvtZPiW5Ync0WjxOd8Hdsx5YqtdYNqiimd7P7WOTwZh12UKPwApUa43Lqy4NOD2asbO2jlcYllOEz9nbc/HFB1ZO3iA0QHicx131s9xwtd22M/qGIHmyjP8AUUhOP1WPim11Ay+Fa3SRs8hyeUMpyf6/dXpDTyJtPVX3UjBA8sUAXHCdpf61Y6bcos0kIdVgfOJFVuZWYg7KB18+leXbFR3Genc+INogdw/wdYaf7OY9Vv15ZdQOQ8o5eWAH3QR197rjvS21loDeyLbRlI1OA7D3j6/w/CmD7UeOtTutRfSNBVHtbQ+CbxlCqzjYiMdAB02yaRd0bg3DNfXvO+ctzudvl1/pRNKrv53PcT1LqgCKOpav9DTBurtUH7iHJ/Cokt1pq8zQWBZl3DS7f9BUGSW1twFMju56RxqFJ+J6j+tQ5JJri5EXKEQH6i9M/nWqomSzzhqOoXN1IEMmVTYBRhR6AVXyylZOUfZ2qRLGYYudu7YFQnX3m880URViScwi0+UX0aguUdfd5hv8iKbHA148SraTSBnixhh9pexpIaXeCxv45JObwSQJABk8vfA86aCa/wAO6brMYtZrsCS4AhMygZt2+qzYOM+npSd6HoTU0lwB3E9T0hpE6+EjA9qKY9UCRBS1KXSdXeEKjk+QojF3NcR5iBJxWWcjqej3hhzC76bHcXQUHPc1V6/w/aa7HyzHYjDKd1ceTDvQ6ut2ehW6TapdpA0xyXkOB8K4XvtR0O2hZbG4jvJsdQw5F+ff5VKgmDALNhZYpwXBbWXhWZFuR0ES4X7q7fRdN4eto2vpViL9Obdm+VC9pqvGHFDRy6Xb3M0co50dCIo8ZxnmJA61Fm4e1Oe8lfiTiG0s5on8MwpKbq4bGOgXYDeiCtuzD7agcM2T7DkwzbjLh6CPH0l2x+7GaFb3j2DVdbj0Phuylvb6Q+8WHKkC92c/lQpqun2RgeCH6XGolDCaSTEjIB9XA2XJ3+VGns+0uCwge9W2SIy/VwNyPM+dcQAPeVsqCruAI+8NrTRlhtXnmCmUqAxAxk+dUM+Fu28hRbPdA2RAPWg69YKXcmggcxbdxiAHtX1SO34Bv4yw55+SFR5ksD/QGvN9vGJHAzgdSfIUxfa9r30rVbfRYXytv+ul/mP1R8h/Wl7byBU5EG56t6/6Vs6ddtf3nndW4e7A9JaCLmtAVXAOwHkO1XsUIiEkjDZFA+ZqFaxc1pbr5sv3VdyCNLVGdgQzGQ/AdKmw4EipcnMNOBXVbS6j5N5FyrefL1AoinnWGMRdD9Zv7ULcHSh9NYq2JUfnX49cfAjI+QqyvFkknZwThtxRNNYB5T95OpBA3D7TS8vV5SBVA0BuZsnpVg9s5JzXJp47Vd66612HAmXjM6LYRpFkiqTULdPs1IudbDnw1YfKobT+LvnNZpLA5MnEpmhbnIr48TouTVjgeLkivsqqY8miDkZk4kfSbWS9v1iGcZphPw8fofKqDOO1DHD8kVpqkUrDKmmMmoc8bNGu2KRudt3E1tJWm07u4NabYpAhDEZXrU67mtliXcZqt1K4NpbtITu29B9zrjsSOY1y1l+Ydr0p8sJLm4h8ckY6VXvPBliAMmheTU5ZJNjtXT6UxXYmtFNGxEz31QJ4hnp0ltkZxk0WWskAQBCM0oI9TkgbIzRToWsSXDjmJoN+jZBkx/S6tT5Ywgc1lc4W5oVJrKzCs18xe6CgaQc3TNMWydoY1ZNiKX2gjEgplaVZ/SeRTsK+k1kDueLAO3iW0fEV+sXIrnFQmu5LqUtI2TVzdaTHBalowMgUKvIy3JVDvXeRgSghgHUjcZbScrW5HIOlCmtRMIiAO1FsNvcSWwJI3qr1i1YWxJXfFIu4BxGiuVzEhqUUf6djM4zGpyR547U1eE52az0uRsB3RpGx0yTj/Sltr8D/AKRHIhLE4AHUmmPwvH9B0m1RTHLLF+rd58kRudyqL3x3JpG9huEppAcmPLR53eBRzYX/AJ6VWcY65baTBPF48dm7xcl1fyEjwIz9hR1yfIbmq+44q07g/h6K7mdrzUbgctvGo96Rv4B2HrSo4htOJuINPuNf1NwBEeaOHrHGSce7+8R3PavM3EWvgcL/AH+09IMqPrAbjDiQS3T2+kwva2nJyIz7TOvmcfUU9eUfPNL45jHiE5f7A8j+8aI9btWt2UMSzuSWZjktjuaphbGaQKSEULzu56IPM/8APWtSgBV4mLqCWbmcLONzzXDAk55Uz3buflVpbQCK+XnG3u/jXa0gWSaMBCkWORFPUfH1Peut+PAlQDqQu/w2poRTEo9UXkHJ2WVhUCdAESQdGGD8al3zGSeUE+6z5Hoa3+jtLpgmXdFPLIO6HzqwgiJVFcfKtCTnbapaIeblPVe3mK5SReG+Ox3BqZSPr2b6/HruipBcuDdW2I5Mnc+TfMD7wab+kypbsqyb46Hzrxto2t3nD2qpf2T4ddnQ/VkXyNegeGuO7PXbNHhkw4A542PvIfI/3rM1FJB3Dqb+j1QddjdiOS507TNSCtNbRTJjdHUMPuNLvX/Zvp9vcNe6Np1umTzGEIMfLyom0fVOYgM+3xokilSUcuc5pEErNZLCjBliz0021kY0vtO5FQjmjZCB8B2q2bWrIc8WmaakJfvFHg0aPo0Vyd9ga723CdiHDsrOfU7VYEx78fxnbzF5a8NtqNyJrwARA58FftfE/lRlDZ/RowOXAA2AonNhBaRYjiVR6Cqy65T02rj7RGy57TuYyku7rwouUmgrXtZFtA/Kcu2yr5mrTifUkswRzZbyFLqQz3s7XE5JJ6DyFXAizN7RP8awv/iAzuSWmXmYnuc71UW8YBO2y4ov4/teU29wB9VuU/A0JW7DLsxwvNWpW2UEwLlxcZerKYbaOMHLAAD7v9akarOY1t4cnAiJPruKr0l8WcE4xzH+tddYYfTrYsfdMYB+ByDVTyRLjhTiG/BN4Bdz2hOPGT3T5EdD9+KNIWEwJYe8Dh18jSo4cumtdRt5C26nkb/n4Uz+ZvEFzF7xI98D7Q/vU18HPt/aFPmGPf8AvJM8MapnzFAHEcrxswQ4FH7yrLbMy77cw9aXnETK8gUnlEn1G6g/HyrRUgiZtqYgyk7g5JJq1srpW2J3qmCknGNxtXeJZUYMFNUeoOIEcQm5FkXI61xeCQqQc4rlYTux5W2xV0OXwTkdqyLazW2IYANC3hDgt722junUsWGRVnqenTaTe+CwwjDvTN9lrWU3Ddu5KluQA1T+1mK1gS3u4iAQ2Gx5Vhi9mu2GejWlUpDCKPW4jc2zsuwA2pb3CEOR60ypLyGa1kGRjFBl3aKWZhWvVxxMXVqCQRB8IQa7BuUda+zJyNUdiTtXoKiAszMczWWTei/hOFnKnloOEeWBNHXDl3FbKuSNqX1WSmBHtJgPzGRFHyxqMdqyqF+IYV28QVlYPgP7Tf8AHT3lboGhay8q8mmztv2Wmjp+nappsKy3WnzRKO7CnDwxollAq4jUN8Ku9Z0+D6I/MgIxWjV8bdzjYMRI6GtDs3HMQ2p6tm0ZVOCRiha3ZDegu2cmrTiBQupXcUP1FcgYoYjWdboEA161ANmR6zGsLb8e0bPDWlx6rexWqthcZY+lN+24L0kWHIbWNtu65pD8H63LpGrxXE6t4JHK3oKeEHHOjLZ87X0QGO7YrxHxcOLfNnbibtBdqx4fcCdU9mOjXOspctaxp4T8ykL37Vw1rhDh/R4LXS7W2Q39yWkEh38CNfell/Ieprtf+0vRhqcCG9i95jIo5tyoOM/Ch7jXjLTve1eznWeO7024sedTsj4DgfPpXl6vH385+nc1scen1gTwjw5ccfcf3F3K7NFAvPM6na0t84jhT/1HHU9smm1xZo1gnDt7pot1iH0cxoUGFGBsi/Db7jQz7F9Zg0b2crMtl4l5qE7XEshOObHuqPgBRBrOq3erWjLIyqEfnRAMDJ61q2Us759B1AVkjk9GeUNX0sz6lcRGMEw7kscKg23Y9h+JNDl7ZgXEGnwZPMBJIxGC5PTPkAOg7U1+JdGjtri8ubjmlg5/ciG3iP2X1Pc+Qpao0tzrIVEDXErEO46BR2HpWlp7M/lE9TXz959mhjtbIsh99dx6Yqj1O9aeUEqB329asNSuPDikQHJJKgUPXchQNn6wGPnWgpzM1wAZHU+L4me5xUmwuWiYlQCHHKyno3oaiwKeRsdsGukK4umTOA5yvoe39qJAfWdvCgucvZtySxnLROenwPlUC5yrMroy8rbqeoBqROj29+s8Z5ScN6eoqxmS3uYVJYIVPKC32D+6f4T51bMptg8UZnVQC2emB1pmcA+ynjjXNcV7aKXQooVDyXd0hXAPQBOrE+W1TfY7w5HecWzate2/Nb6QOZFcZBnb6vxwMt8hXqSyvGsdPJxzXEx5jnt5UjqNQVOxZo6TSBwLGizu+HeL+FYkllRNZt1+tLaIVdfUxknb4GpWj8cWUtwtu8wjmGxjk91h8jvTInSM2glubhjK25IPT0FDsui6bqdz4eoWVrefum4iw3ybrWbvz803AuOpfWOqxTRKecb1dpqcaQYDCgl+CZrZA2iXU9hj/Zu3jRH5HcfKqjUbbjezTEKWN0B3DtGfuNSOJbyt3Du+1uPkI5xtQbqfE8EAZjJk9gOpNLfWNb4qtNQW01K2S08Qcysj84YehrW3V5mEkrl2PcmrbT2ZQuPlWS7yebUbo3E5yScgeVbpbjk6YrZIxkd6nLF+r6VMHFtxtp/0nT5Y+XfGR8e1KKOQiFk+0N/jg16B1+1MqHakZq1gbLWJ4R7oLc6fA1oaZsjbMfWpghxOlrJmR/jzD4Gp2rAvY2tznIGY2Px3FVFnLyzqrjBxykedX9oIrzTprF2wQeXJ+yeqmiNwcwCHcuJGsLo5STmwcjPx7H502tD1BbqyibPvAYI/5/52pHRSSWt09tcDkIJU5+yf7UccL6lJHdrA7EAkDc9M/wCv9a7lG3CErYMNpjMCYLcueRvLsf7UGapaM6t4a5QEyFT9gr1I/CjrTmMgWRVzjZx5H/pVJfpaW1tfTXQYRpMseV6gMCG/DB+VFY7MhfpiTtDAFvrn8oHaBw/d6rdJGibuevxp0aV7E1urZWkunDkZIHShLhq6g0XV445mUhCBzDoR2PzFeidC4o0o2av9IQbedZes1moUgVdRnTaWkqS/JinufYstmGkS5kyvnQpqXCb2CMWckCnxrnGmkRRSD6UhOOmaSfEfF9pOZI43Vh6UKptZcfOD+ktdXpUHGJM9nOsXOm3MtozkwBth5VZ+06c3+kAxSMO+KU+n8VHTtTeRd0Y5zVhr/G6X1n4atzHHSufQ3C4Ptg01dXgFMztw7pa348GRmHrmrG/4ZgWYQwyHHck0JaJxWthJ7+2epqyk4yja7DBuYE9qIadRv4Eor6dq/N3Cyy9mNjfQhyzMx9amR+x6xLe9GT86n8Ncc6VFEiPJlz2plafxFp88IkrNt1GrrOCTNOvT6VxlQDFxD7INMUDNuKlJ7LNOjGFgxTM/Ttl2ArU65a/uilTrNQfUw40tI6URbH2X2JP7GsplDWrY/YFZVfxWo/qMt+Hq/pEi6PxRIkgzIasda4onlsWUTHBHalRp9+FmGWq7mn8eLlB616rS6FQ4PpErLwUz6ymnMcskjHcsa0tbOOW8QYGM1arpMXIebPN1qDGv0W9xnYedepaxSMLMMVnOWhMbCFbXAjHShPXIooLCVnODjIHp50ZRXEkthlItiPrE4pfcTXdxG0scqHDfdWTu3nbHWAUZip126uW4uh8IHmMMccIG4Zd+np1ozZP0jwWba0ywhHixjzK/W+ZGaDjcypc+EltJde80SeDtJFz9QrHYZ8qYXCtvDCFWa6EQGAIGAZ/gcbCkbUG3afSU07HeT7y/9mOtCXguLTgMTWLtA/oCcqfmKZFpCZ08PpkZJP2B60pYLObgrjFNQWJn0m+X3wo38POSQP30O+O60wdY4p0/T9MCaVPFISvM1xIf1aZ+0e7H0FYt2oVV8vZm1XUScH0g97VbvS9P0a3tI1AlB7fWRD9Zvi31R8aXHCPDE+pG/wBdkRY/EDRQ42UPjoP4Y16nzqPd+PxBrjXetXskNirll8RsTT+fKv2c+f2R61Ya37Q0sNJ/R2gJHByp4QnVfdhQfZjB6+ZY9TvSyLYF8Ovkns+0uxTdvfofvF1xLpyadqckLZDp9WM9VHZmHbPUD1oLuTzzYz7idT5mrC+1R7yeRFd3Zm5mfPMznuSe59arDkuF2wPsjfFeiqUqo3Tzt7qzHb1JFsDyPzfuiuckhiljlABAx171KQBIpCeo2+4f3qDcg5jjznlG/wATRByYFuBJN44ZxynbqPga1ikd+RVUyeMvhsg35j0Hz6VzuMCVRn6vKv4UyvY7wkdY1063dxZs9Pf9SCNnm7H4KN/jiodwi7jLVobH2iOT2f8ACx4e4Xs7KYZlUeLMf3pT1Hy2Hyo8ihJ99zvWttCqIoxhV6D86vdPs/EYSSDYdBWISWOTPRKAgCjoTbSNIEswubsZA+qp7VL1KxtGBBGwqweZYYuVdsfhQRxTr0dlayt9IWOVVLjmOAQOtdtxOVixlmIYVXEdxIuPJ6jyw3h/ZXhYeTgNSJv/AGnamuru2mxeNalRkseX3u+PSrTS/apqjyww3VkkKySBPGLZCjvUmtsQgYe8KuLtHOrxPE8apeQfrIuXYN5j50B2xKjlwQRsdqa95jUbBLyBwZVHMCP6UutRs5INVNxA/JHP7xXsD3qy+0lgOzPtupY5q0SLKb1HtUJPM27HrgYqzC4TapMGfpB7UbcOeX0pScd6TyKt7GnvxHf1HenReJncUI8RWC3VhICucqaNU21gYregdSIiLiM8qzRnB65FdLPUGiulkI6jldezD+9TJLdozJbsP2bFflVdLbDP51rlQwnngSplzdW8GoLmSURyr7qzkbMOwf8AvWtol/pdzF9ITEYPuyqeZfvqst7m6tT7jZHTB328vWpI1Z4RhVeJW6qu6n5GglGHHpGBYhO71noPTJXt0S5YALPEoYeuM83w6jNC3GF1J9Bs7GKMma8drhlA3JY8qLQBpvHGtWIRLZjIijAWVQQB5b9qvZuNJNcv5ry+eKzvnjEcUixZSLbBKgdDjYHtVSG3hscRgWVtUUzzOOo300eryBDlIQsIYdDyjBo74av5ri3QGRsH1pfQaQrgiLXNPbI3WSQofxFGXDcP0GFVub6x2P2LgNmn6DUpG4xGwWEEgdy04kUDBB6il/dH3jTF1eAX0YFrPBM2PqrKM0HXXD+sAkjTpXH8OG/OtQ31EY3D9RMzwbc52n9DBiVds1DkyDjNXlzo+rRLzSaXdqPMxE/0qnniljJ8SGSP+ZCPyoTMp6MsFYdicFGTU62jBYbVCjK5+sD86s7QAkEYNAIh1l9ZFYCrjqKK7bi76NAI/Exig1mKw7UO31w4fZyKytVSLDzNCq81Dyxw2nF7TOAJc/OiC31x3Ue8Tmkjw/IzSAliabXCdut/qsMDbjyrHuoRBmaunvazuEaX1465SNyPQVlO7ReErVrBC0KnbyrKyvEz0sfLqDgtPMsF0yyjeiK1vnwN96EEkVZAc0T6RALqRRzYr6LQo7M887EjAhGNZVbfLJlsYoW1DUm+k842yaLptPSK3xgHahK90qSWSWULiOMZJ/Kjo1eYF1sxLyLiGMWaIHAwuKHru8+mXnhzsZo2OPDT6z98Dy+NVstgYwGkdgD0Vep/tRPw9pmj6YIdQ1uKWe+uwVsNNgy0soOxYjrv2JwKztSa9Ohf36jNJsucKYu7fh7U+LLmMW1i5tUlMxZP1VvCOgye+3c5NHWn2/D/AAjHz3t4l5cKNkg91EP8xod9onEvEEN3+iLi4i0CzgUctlZuOZB2DMO/oKWrPEsLXrWj3S55RJdyMxdvJR3rCNl165Xyr9Of3/xHlWqhvN5j9f8AEa2ve0fTdUsH0428Tx83Oi2oLyRsOjB+xoBfXtVH6vxpFhU/qwyAlB/QVQzajd+GEZkgB6RQgKfwqtkvJUP6tcsf3iSPu71FelUd8zrNWScjiXF3q1vG/iXEyq7dSSXdqpby/t7nAS3muGf6viNyr/7RUD6RHNdBHs4p5Cd+QFfxzVmI0DNIFKs22e4HpinAgSJGxrPWVUkcmCGCog6rGMD4etcYFCsZCvf3R5mrCcoXEZPu+S/0FYIwiGaUKuNlXsPSr7uILZkyNI/hxjm35feb1PYVCj95jK++Dk+p7CpLo1wxJPLGu5Y/89a5SlY1HKOXA91fL1PrVhKNJGmaXfa/rtpo+np4l1cy8i+QJ6sfQDJ+VeyeFOGLPh3QbXTLRf1VugXmIwXPUsfUnelf7EOCH0/TG4o1GHlur5eW3VhvHD5+hY/gPWn/AGNo0hBI27Vm6mzc20dCa2jq2LvbszLOwMjh2G3lV/DCI1GNqyKNIUAPWud1epDEcEUFF94w7noSv1a8W3jYZ3rzz7RdXXVdXh05DzJbEu/qx6D7v60zeL+IFtrZyrZc7KPM15r4k11otTnsLZj9KJ5rifuud8D18z2oqruPEgMK13NLhY4VPKSGcfZHb41u8JdCrDIIxiqDSpThd6LLaMyKNs1JGJHiF4fezq6vpNN8K7uTJGuQqkb7HG5+Vd9TKTyTFMYguWT5EVH4MaKCzmVjho2OR6HeoNrJKtxqNteNySu5cZO3NnP50HHOY+RlR9pY243GKnsSE61W2rHbNTpT7tdF5ElUOpqmvIPEhdMdqI44eaLOKr7iEgkY3NXWCaIbiSyNnrYwMCaPPzBxVBIuCVpke0SwEAsbsLjllaMn4jP5Uvp0w2a2aTurE8/qF22kSFyb108FXTlI+FfeWuqDHrRYvI6KTlW2ZetblQo2OT5VvOhIEqD3l6jzFaoodQw6GuxJkqGcungzklR9V+6f3HpX2K6uIT7r7DuK5BAK2xtiuxOzJo1W42yxrY6tcZPvEHtioBFfMda7bJ3GWia/qUX7O7nT+WRh+dSV4u1uM7X8zA9RI3P/APLNUeNqzlquxfaW8Rx0YT6XxELnV1n1i3tJ4beGWdQ1ugLOq+4NgM79jUO11e4vL1vpYjfn3JWNVKn05QNvSqRdpx/Ka7WnuXYNQECnIk+IW4MLGAZcVWz6ekr5K0wvZ5wnDxZcTJOzckBVWVTjrTjX2F8OSW4Kxzq+OviGszVfE6an8Ns5j9WhsdN/GJ5s0u2jtsHGKN+HNaj03WYJwQADg0x//wAO6RDdmCRJGAPXnNXlv7EuGZUDCOUH0kNZ1nxHTOMc/pHU0lteDxDzhvjSzk0xC86jbzrKE4/Y/psK8kNxdqvkJjWVl+NSOAx/SOGvcckDP3nmGLUy8oAOd6OtD1Lw+UlsUtNFhM1xlhnBpk2umk2fMqEHFfU1CrwZ45GsPIhI/EBKkM4xUjTtSjuLOeNlDeJkj5UCy206S8rZqdDK1ngtzFEUsVH2vIfM1TUUDZ5e4ejUuX83UvOS3jVtS1FkW2iBcqxxzgbnPko/HoK24e4ztbSDUNduSkWqXsbSPeuvN9Atukcca/vt5dhQFxFfzXNta6Rz/r9SkzM/7sSnp8OtDGv6kHt49Ns2EaSzgu3bA2HyAFeY1e660oehx/n+fSbNRWqvdj6/4k26U6hbTa7M7J9KuGFvG555JQD70jsfL+vSqG5ilmuIoo2Ow7bhR2xUi3u8WKLuy5YRDzXOw+Zog0zS4vASedgrud3PQDufyFGYpUoEWUNa0EmtAjsMc3KcEk9T6nyqovmKuwDgu/cDp/b0or1VV3FqhCMSEyNz8qGp7VUuI439+Zm97fZf9ahTnmUcY4EteHdFT9HG8ljy8rcqAj7I/KumoW7FzFbKJJO57L6ny+FSG1Qcsdja4UABAQPqr/c12N5bpEbe05C6/Xkb3lU//Zqqc5zCjaBgQbktEsvekYvK3fufgKiNFIxD3HuL9lR1+X96n6hqFpaSMEzLcHqT7zff0HyofkuZrlnkc+6vUDz/ADogUmAZwOBOl1eKMRxKvu9AOg/vRr7K+BpONeKBJeRk6XYlZLkn/aN9mPPrjf0BoDt7eWaWKGCMyTzMEjQDdiTgAfOvaXsw4KbhzhSDS40AlRTNdzech+tv6dB8KDqLBWuB2YXS1G18t0IXadp8QKoFAUYAVRgAUSLFHbxjoDWllZpBAZX2271XXd8OZsNsKyxNlvpOlxeYYgGhHXNWMSPzMQRW19q0cau/iDbrv0pU8ZcZWFjbkXd2IxJ0Ubsw8lHf+lEXJOBIKhRuaVfFfEMUNpc6tctmGAYiQn9ox6D50ibW5mu9Rmurhy0k7lnPqTmpnE3Edzr96pZTDZxbQw56fxHzJqt04lbsI31W2NaddexTnuY11/iWADoQzsOaMqe1HGkSq4UZ3oQ0+ItGCeo91qINN5oLhcdM0o4zHajiGaQyDeGRo2IwWXuPKpiRySSNJKq87HfkGB6V9swHhRj3q6t7YEA0vHwxxiQ4IiHAqTdRFVXA3NSXt/DIOMVNvIFxAe2M1EiR7aEmHBHQVAvIirh+XpVp9LghwuQMVUaheoQeUZq4MoRAH2lwrJwdJOmCYp45M+W5H50opPejBG1N3jITXHCmqIYz4fgl+ncEGlHCOaBcjtWtpTlCJh60YsB+ki4Oa3UV9dMHpWLsKaiM6AVwGIZ+Uj9W529D5V3BONq+SIssZRuh/CpnT7j0rDttXOB2PNFJ+0Tv5jzrtg4rp0+YOKwDatvQ1993O9dOmgG3Ssx6Zrpivnfyrp04j/vGMfY/OusO1ypx3rkM/SW9F/OuiH9eDUHqSO46PYtxDDovG7W10cQ6hD4a/wD7FOR94yK9aWmuae1sHU52rwJp08q6pafRyRMsychHXm5tq9LxcQGJGQPXkfi2nPiixOyP7T0mgYWVlG9IcazxPY2t2zvtk1Y6RxXYSRDB60heIdXkmcnJ61L0DUZQq5c1lnRsE3A8x/xELbCOJ6MXiGxI6VlJxNXcDdqygeFd7y/hUxAcOmKO4/WY600LTUbaO0KkjpSThuXjkBU4NXdvqF26hec4r68VDGeBS4oMCMGa8tnYEkZFbO1pcIeicwxQQrXLH6xNWNrJKgaWTLLGMhT0Ldh+dQ7ELJquYvjHcq9WZRxhHGT+zsm5T6460uL68zMcHYZOfjR/rLOdcsL6UD31aFiBgGgG+sJFlOwVVYhi3Qb159l2uQ012YsuV6hbw3ZrqWo20cq4jVFyB5f69KNtWawXxI2le3hhBihEagsz/ac57DoKCOFr1bK8ieJvEHMGZ3GOb5dgKja3qk9zK0SsNydycD6x71nurWXfQR2tlrp+pknVNYhupGFhZR2scEYXmB6KPtOx7n8egoRa5Y3XOTnqR5n1+dbXV7ywixiPOiNzHIx4j+ZHp2FQFcfSOUnnc/Wc+fkKfRNq4Ez7LCzZMlC9eNWdT+tlPKp8vM1tc3rWsC2sBxJjJb93P5mocSFr1A3RTXONHmvGkbcM+auBBFjMjiJLEjLD3VHmxqY9qlva4f6oOG/i8/xqVpsC4a6lHupzPXydHkht2P1QSxHqdxU5nAQ09jWjLqntIjvp4wY9Nha5CkbB/qp9xOflXsTTrpIdPFrGuOcgufPHQV529henFNG1PVmX3ru5ESHH2Yxv/wATfhT+sEPMuax9S2bDN7SJtqH15lvqV2YrPC7bUC3moPHHIxNFeqtmIjPQUA6u2IyM9aCvcOeoqPaXxmdEit7W0YPqc45yp3VY/Nh6np86Rl/eXuo35ur+dp7h8ZZuw8gOw9Ku+NtSbVePdSn5uZI5vAT0VPd/qCaHGYF3k+QrYprCrn1mDqLmsYgngT5yhs+n9K725VZAzdQd/UedReYjDCpCKrlSNs9CelHMWXuG2myTc4VTzBgCDRXptrcSOM0J6DG0trEQo8RG5QCev8P3f0ppaEsTKoYA5GQT5VmWcTbpG4CWunRyJAFkHSiGzlAABxWgt0EAAxt51oImUjApfMdxJ1w4YDH3V1fxLorGuwUYJrhbxMzjmFX1lAiOG5R86qZYDEpf0WCelbfoeM/WWiOZFyWAFQ5ZCoz5VIEgmBHGVhFFwbrCKo5jZy42/hz+Vea4D+rUrscDbqK9ScQI13YXFudxLG8ePipH515ciTkULjddjWto+iJh/EBypnx8k4Ycp/A1zIwcVIYA+ormV8veH40/MufB2rcfCtOZEHMzBVHXm2qK89xcHltUKp/5jbfdXTpvcyRxyIwb9aD0Hl61JB5gCB8zUaCzSI8zHnk7sakgnua4CdPiyBndc+8hwa3B3qLCf8xc/wA/5VJTc7V06dDtvWhO9azShTy1rkla7M6fI8GeQ+QFbZAlyNq5w55pT6gfhXMsefAqD1JEt9P1oaLqtjqZtIbwwTq4t588koHUHG+MdxXrjheb2fcXcN22u6bY2drHNlXt5y7PBIPrI3vb47HuK8TPIbi7GDlY/dH5mmb7OdVlsb66sFkKx3CCQLnbnXv8cZFJ3adbcE9xzT3ms49J6rj4I4HuRzTWemt8RIP/ALVPh4D4Jh2hs7FTjOFmmX/7Ukk4huI+krbetbDim+AblnbLnc57UL8HGfHBOcx3PwZwsDtbWw+F9KKykpFxRemctJOxAXAGfWsqp0Q9hLjUfUxIwJzyDNEVpCqKCap0tJoG5j0qzhaRonYsAEXJr1PRnmAQJe2xQEdKmMqFTEBtjmz60Kx3MgYHnH31fWd4LnlRBmUeR2+JPapsHlzCU2DO3E11XTJLnQppI1y9swlH5igfVYsxG7VcxuMr/Mevzpmi4ilj8COQMiE/V35z3OPwFC0OjGTxreSPx0hnbw4AcGQkZCnyAG5Neb1hKv4jes39OAy+GPSCOm/5TTWupTyh2xEnd8dx6A9+9VVzMzZcjq+M+WP+c0Qa3bvFqaQlxI4j8RmUYX0wOw8qGbpXFuqn7THf+oqlXm83vKXeXy+0gc2FLKTlj1PYV1ghZgWA9K+wRiUMAMhTV1Z2qpCwlHUZUedMRQCRIoOYLN8jUuy08+K6suDHuRUy3azSVY5Tyxy+62ex7GrGSMQzJyuJOVfDkI2Yr2OO+Kr0YYDIlO0Xh6ZOmNxzIf61HkuEW2C4zzxq6/zDapdzmNpEJ2Ox+PY1twppf6e4u0fR9wGuQsnpGPeJ+4VQtjmWC5IA9Z6d4A0QaJwbpVgVw6QK8n87e8f6/hTDswAQao7RQNwMA748quoDhc1inLHJnoQNowJy1OUcrUvddm5FLZ6ZNGuoycytvS14quPBsbl8/Vic/wDCauglHOJ5TllM97NcE+9JI7n5kmuDjIUjpjHzrbHLHnG7KMGtkAPunoa3J5k8znGNsdq7xgxsUk+oe/l61iRkZDbEHFSRHzKqj62MgenlUEyVEI+GJQtxLZzbo6g9fLypl6TdiIKZmHuOFZ/U/Vf4HoaT+k3P0fVLctsFYfd0pp6Oql57eUcyD3D6of7bUlcvOZraZ/LiM6zkZ1AO9WkcKsBQvw7O50+PxG5pIiYnz3Knl/L8aLrchlBHekjNEGd4oQCCKlplelck37V2zgVXEmZJIcb1AnkBGKkyNtVdICcmrZkSDfsux8iK8w6jAbbWr+1/8q4kT/iNembwZjOPKvPPGEIg451ZQMBpRIP94A1o6NvMRMr4gvlB+soj6V87nNfW6+da7gVpzFmMiMo51Vu+4rYY6V85vnWyiI5zJyHt7uamdNTtmtevauvIgB/XKfkcmta6dIMP/eLn/wDZ+VTohk9MVChObi6//ZUyE7daqJ0hSsWvCvlUpQQm4xUKU8uot5VLd+WIkeWa4Tp8iI8Bm82NRLibkQ8v1jsK7BuSyj5tvdzUOCNrmcuw90dKg+06SLODkQMRvVxp93LYzrdwnDo45flvUI4RPhWglJCgdFqCJZTg5jmimW4t47mL9nKodfQGugJwaHeEb76To7WhOWtmyP5W/sf60QE9qKDkQs3zjpWVpkVlTIlE/H1kTleD9FHxjc/nXWL2jRRqVXhPQgD/AOg396W7nl37ZrQyHJya4gfwmBDH+ARqWntPgjYRXXBWjXZJwvgoYmJPQd81Y6j7UtFsr2G2XhG2dIRmaKFwsfieRP2sfdmk/b3TW9ws8YBdAeUn7JxjPxFaxQGUO7ScqIMu7b4/uTSz0qTn/wCmM13sowP7D/Ee2ie2fR2vhHBwcsDEFi0ZQcoHcnsKgxTWut8Q6zq+nupkmkEJRBgRyYywHpv1pT2sUsi21taxBEuJwzjI5nSPclj5Z7dKddlZW/AHszl1q8RW1C653t0PWSeTofgoxWRrmCAKvZ4Ampotzks3Q7iw1yGL/EGqH6QFUMlnEfIge8fkM0HatOjShIE5YlHLGM5wPP4nrRrJw/Lp+jwapr8vLNchngtmOJJsnJcj7K+ZPbYUErBPquqpbWyeKzMdl2yO7egpnTMNvByBFtUDnrucbZWg5HwSzbhR3X/WrqWYLBGqKQykkE9QKKLjheLRrWK71KeJJphlI+rkeSr1wPM4FC960BnZo1LDsXP5Cji1X5WCNTVjDSpvZ3cBEX41vbapK8a2lzguv7Ns7/A1wupiMxgjm7gbYqKIORfFY8oG/rV/TmBzzxLeS5cpzsfEQjB5uoFNL2GaB9N4g1LiEjMdlCIEY/vv1/4QfvpW6fEL+9t7VT7s7jm9B9qnX7PeLtC4T0+64emga0iurozx3RbK7gAI3ljsem9L3K5rO0RzTlRaCxjut/rdRVnz8sWaHdOvY5gHRwyncEGraSbMfXaskCbZMgahKOVqVPHlzycPam+cEW0n/wATTIvpMo2DSi9oc3Lwxqm/WAr9+1HrHIi9xwpnnrJT3GG3cGu0cXNvCcn909f9azkVvdY7dm8q+rG8Rwwx3z2PqK18zz+JLhjSciOQFJOxx19MV1MWW5VOHU5x0PxHnWod/DDo2SNip3rcst66jOZAMoT3/hPrVOYbjE2MeW5yMMpHN5j1+B/CmfoUwkWOcbiSLf4gUsrdzNywyHEq/spD3/hNG/C90Vsp4iPejUlB6nbH30vcOI1pm80ZmgSqwuGUe68pb45AowsXI/Vduqn0oH4f/V26oT/rRlaHLR798bfCkTNYdS+Q+7X0nvWkYJUZ39a6EbVWWkd9+tRXG3SpLntXA71E6V9xGSh260hPaRAYONy4AAnto3Hrj3T/AEr0HOPcNJD2r2/JrmlXH78MkZ+T5/OndIcWCIa4ZqMXbDzrTOa6yDvmuXetqeemHr0r6POs7VsOw7V06fPnWdSK22yT0r536YqJ0r7U801wcYy9TYD1FQbIZSVvNjU2AbnaoE6Qb33b1G7Fa6TuRaHfqMVrqQIMb+RxXKVh4CD+IVE6bXTM7R26dcAVNjjW3hA7jrUa1A53uH79PhW6y/SJyV/Zp1Pma6dMlYtiMdepr6BgdK3jTq56mvjH3uVetdJhFwlfC11qJXbEcuYm+fT8aYrEA79qTsEhhmDg4wQc/Om0JllxIjhlYBsg56jNWU+kKORJHNWVzDCsq86KqUKhPMvOjDBXOP8Aoa6SaXcXjTXWn4nhjiEhXPvgDYjHcjr8N64M3iRBu+K1t7q4s7gT20zwyr0ZDg1FgJ5XuCQr03UjxBWO74HnjNTFhluoo7S3yCX5gACzSt0GAP6VK/TKXDhrzQ7G6kY48RQ0LMf904J+VHGncSWXCloAmn2dnqkg2S3QyypnplmOc+gxSV17oMBcn+fzqN0UI5yWwP5/O5ccJ8CWujG31/jSVLdYlXwLCRxzNjcGT578g6965cWcc2t9xD9Pa2F69sOW3W42hh9Qn2m+NU2o6rFHm5v5JLrUZvsO2WU+TN2x3C/DNCt5cDwjK0YEjEhGxhF88CsdaGufxbTn09h+X85mw1qUp4dfH9/zkfiLW77VLiS8vZTLNMdzIfeYdh6AeVFHCMthwVwve8VX1tHda9qCeDpkEoysER+tOw9Tso70CsYpZ7aFvfkaQNJk747LXe+vru7Y3l25Z5JCR5DAwij0AHStE1AoKxwP5xMwWnebD3LTU9Vlv7xXEnjXMiB7ieRuY8x35T54+7tUG6hlii5znxnxjPVfU/kK68OtYWVpNqt+vjyIcW9uejv++/8ACPLuaNOFzo9xczX2rJ4yW8Zky3+1mO5+Q6fGrMwrHA4ElFNp5PJi7/REkMS3E6MindVbYv8A6VS3U0kshAPuDf0+NMDi67inLtdSmGecZjgQfVXtzeQ/E0Iwaes4lCSRsVwWDNy5HpmprfI3NB21gNsSXfBVp/lLzVZAeSMeBFnzO7H7tvnXbUZxICMjFXEkUWkcKWunoOV+Xnf1Ztz+VDM55gN809UOMxW04wsLODuP9V4XkSCUtfabneBm96MfwMf6Hb4V6C0TivRuItPFzpd4suB78Z2eM+TL2ryUSa62Wo32m3iXmn3UltOnR4zg/A+Y9DQbtIr+YcGMUax6/K3Inq3UbgCJt6TPtMvuXhq6QHBlZI/+LP5Vz0z2pTTQi21623xj6Tbj/wCSf2+6hr2iara3uj2otbpJkkm5soc7AHr5UgtLI4DCP2ahHrJUxeB3Dcv1WBzU+3lSQeDL9Q9D+4f7VwtlhmhEbnBHRh1FdUiaGTlfrkHI6EedOsuJlo2eJ35TBIVcdNtu4rk36u+9w7HDbVOOHTlIy0ePmtRZUAmj5e64FDEMRPs0ii5I7E8330V8P3PLfgM/UB8Hvg70HsviX5GfdTAPyq0spzBrNlIW5VVwHJ7Kxwc/Koddy4lqn2tuPUeWhMstrHIoO/XPY5outG5WQnoN6HOFOHtWnguEt2t5fDnZOXxeUnGN9/QijCPhriNF/wDDC480lQ/nWeabP6TNhb6yPmEnwXCnAzUh2ATmJ61U/oXiZDldGuTj+X+9WcekcQSxKDpcqnH23UY/GoFT/wBJlvFT+oSK75bbpWvU1cW/CmtSnMpt4Af3pOY/coNWkHB8aDmvNTYjuI4wv4sfyq66W1v+Mo2qqX1gZMeVSTtilF7WF5tN0u68NuRLl4y+NhzKCBn5GvTiaJoNunMlt9IkH2pSZPw2WlB7d7drn2eOUTC2d5DKN+gPMuwGw6in6NIyMGY9TP1GrDoVUTzezAjrmuYwc7185WG1a82BT5mTOgXfJrPx+FPX2CeyLS+OUveJuKpSuiWDciwhuQTMBlix7KPxoh4p4h4At9Ray4U4at2tIDyCYoFV8dx3xUL5jiXK4GTPNJJI2VvuNYVcqTyN0J6Hyp7f4n09R7vD1n93+la/4sth9XQbIfKieHKxBWFvKLbeKTJP7pqbHbzAfsZP/Yad44vRdk0WzUei1seM5R9XS7Qf7td4c7iIS/tLqSBgltMxBBGIzUF7K/eJALK46/8AlmvQrcZ3ZzixtB/u1xbi6+J2t7Uf7ld4X1ncRDNp+pzBYILC45e58I1ZQ6LqUUKomm3RHpE1OleMtTQ+6lsP/wCQrf8Axzrf2XgH/wDIV3hfWTxE0dG1lgFTSbsk/wDpNX0cPa0ikjR7wnv+pNOI8c68elxGPhGK0bjTXmG94B8EFT4YncRPrwzxHM4P6EvQme0JyaJuG+GuI49bhZ9HvYIQjlsxMA22wPzo5/xlr3a/YfACs/xnxB0/SUg+6o8ISQQJz/QmsdtNuP8A2Gsrc8X66xydRl/CsqdsncIkoWzEQe1avW0A61kinqKiBky1cafZHVWw9yWMdqG3Cn7Uny6D1rXSS8LSapKfEuC3JAX39/u588f1rneIWsrRAdolx/7t811jbFlGB0jiYD4nrSbrlSfc/tG6zhwPQDP5+85yXck95HICcKcKSeo8z8TvXeV3eyQE5TO6/PqKglQqRspyAfuqzsopJeaP3fCO7MxwF9c1ZkAXj0kI5ZufWVCpy6uJCdgec/IbV9F2klqYZRsfeUjqprpeIkV8eV+dRkcw+0KywszdDxLlVt7ZSWXszf6UHIxkywBztEkwWf02aC3DMLZQMhPrSHyHl8avp7O6t4hyMLeIDCorDaq063HZxlbKNI06GRhu3wFRZtZv5QOaUgnoMDP+lU2M0MGRB9ZHvvGYHmcuR0Z1yw+dRtHtZL7Xra0kJMfPzuSPsjc1wmubueQp4rOe+DtV/wAORLZyRSSN+tuTnPkg6fef6UdKz1F2cE5kriK/abUGQbAdqqPEJG5rtfkSX0rZ71FNOgYGIqxySZuDnvXzFaGs5sHr8q6ROo5VBY9AM1AVi0RLAEMScEZre6mK25A6t7tceblQL5CqmTNMRq/MIwP5Tit5ZwYgOXBByDzZrkWya4ytzNyjpkA0BuoRe5ZQSlrkPjbkOfvr5M4ict15RlfyrnBIkOz5d26IPzreeOSdVKYJ748vSlscxrPE1tQArO/Qbsf+fWpMf6xHZ/rPnNRmXlhRFOV6n1NdVkAj5F6kbnyFHRfWLsccT0J7NtZnl0iO6D5aQK7b9WwA34qacFnrLso5pT8wDXnn2RXavZ3VmzbwyZA8gd/65p0wFRGKcrXKwgPEM4tRVx+1H/tFdhejG0v/AAig5Zym4au6XhzuaL4ZnQr+mZ/2jH54/pWpuExzcoz50PLfDGM1t9MzsDVvCnS3mu9utLT2pL9L9nmuxncrCJR8VdT/AHoyNxzDrQvxgi3HC2rw4z4lnKMf7pP5VBrlWnlbl5k+O9aGMYye29fVlXwEx15RXfTIm1DWbOxQZM8yR/Inf8KXKwAj3guLzhb2KafotvcvCdTAaZFOMg7n8qAubHQ0Wcd3ok1S106M/q7OALjyNCFFxjiXJ5nTnOetZz1zJxWvNXSs689YX9a4c1ZzV06dues5tq4c1fC9dOkjm2r5zmuHOa+c3rUTp35q+81R+avvN6106d+es564c9ZzV06d+f1rK4B8VlRJgAhKnArfmJ/6VlZQvWUkmcn9Hg9/oyt07h8CtbL3oirbjLjHyrKylj8p+8bHzj7SIhP0Rz3DLip1sxNpdKTkcmcfOsrKJZ8hlavnE5ah7tq0o2cFQD8q4F2OkxuzEs4BYnfO9ZWUoOhGG+YyND793l/e5VBGexrSWRzbk53YnJ86yspgRY9ThOzJiNPdXGcAelELkpqsKqcBY0AHl7orKyiJ2YM9SDIzeO+/2q+cxrKyjwU1JNW1/BDHZTKkSr4cKupA3B5hvmsrKHYeRCIODB26ZuaIZ2zXN2PnWVlQZWaEnNRgzA5z3zWVlCaWWdYXYF2zvg71ZWjt4URzvzY/pWVlCMOkyY45wOgk/KuKsc9fWsrKKnUG/cYXspmlXiW7jDkKYASPUN/qa9AQO3INzWVlO09Sw6nck+ZrUM3N1NZWU4suJ3UnzNSYsnqTWVlWkzqwAqi13/wy69YX/wDgaysqJB6nklSfDAos9m6LL7RdLWRQwBdhnzC1lZWd6xZexDDiZ2bivUcnOJSPlVXmsrKIZaaMa0yaysqJE1yazJrKyqyZ8JOK0YnIrKypkT6Ca+gmsrK6cZmTWZrKyoM6YazO9ZWV0mZk4rKysrp0/9k=" alt="Gabriela, avatar ejecutivo realista de CONTAVERSE IA"/>
      <div>
        <div class="name">Gabriela <span style="color:#f4d281;font-size:12px">· ASISTENTE</span></div>
        <div class="desc">Bienvenido a CONTAVERSE IA. Te acompañaré en el análisis empresarial.</div>
        <div class="note">Imagen de Gabriela integrada · Chat demostrativo disponible más abajo · Voz pendiente.</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

registros_panel = []
if LOG.exists():
    try:
        registros_panel = leer_registros()
    except OSError:
        st.warning("No se pudo leer el historial para los indicadores.")
total_registros = len(registros_panel)
total_aprobados = sum(r["Decisión humana"] == "Aceptar" for r in registros_panel)
total_pendientes = sum(r["Decisión humana"] == "Modificar" for r in registros_panel)
total_rechazados = sum(r["Decisión humana"] == "Rechazar" for r in registros_panel)
st.markdown("### 📊 Indicadores de decisiones registradas")
kpis = [
    ("Decisiones", str(total_registros), "Historial local"),
    ("Aceptadas", str(total_aprobados), "Validación humana"),
    ("Por modificar", str(total_pendientes), "Requieren revisión"),
    ("Rechazadas", str(total_rechazados), "Control de riesgos"),
]
for col, (label, value, detail) in zip(st.columns(4), kpis):
    with col:
        st.markdown(f'<div class="cv-kpi"><div class="label">{label}</div><div class="value">{value}</div><div class="sub">{detail}</div></div>', unsafe_allow_html=True)
st.caption("Los indicadores se calculan sobre el archivo local de decisiones de esta versión; no representan saldos financieros.")

# Chat local de demostración: no usa API ni servicios externos.
def responder_gabriela(pregunta, casos):
    q = pregunta.lower().strip()
    if not q:
        return "Escribe una pregunta para que pueda ayudarte."
    if any(p in q for p in ("hola", "buenas", "buenos días", "buenas tardes")):
        return "¡Hola! Soy Gabriela, tu asistente de demostración. Puedo explicarte conceptos contables, tributarios y de auditoría, o resumir el último análisis cargado."
    if any(p in q for p in ("último caso", "ultimo caso", "mi caso", "resultado", "analizaste", "análisis realizado", "analisis realizado", "resumen del caso")):
        if not casos:
            return "Todavía no hay un caso analizado en esta sesión. Sube un CSV y selecciona 'Analizar caso' para que pueda resumir sus resultados."
        c = casos[-1]
        cont = c["contable"]
        alertas = c["comparacion"]["discrepancias_y_pendientes"]
        return (f"Último caso: {c['descripcion']}. Base: {dinero(cont['base'])}; IVA calculado: {dinero(cont['iva'])}; "
                f"total: {dinero(cont['total'])}. Dictamen: {c['comparacion']['dictamen']} "
                + ("Pendientes: " + "; ".join(alertas) if alertas else "Sin alertas documentales automáticas; requiere revisión humana."))
    if "asiento" in q or "débito" in q or "debito" in q or "crédito" in q or "credito" in q:
        return "Un asiento contable está cuadrado cuando la suma de los débitos coincide con la de los créditos. Eso no confirma por sí solo que el reconocimiento o las cuentas sean correctos."
    # Respuestas temáticas de orientación: no son consultas normativas en tiempo real.
    if any(t in q for t in ("retención", "retencion", "retefuente", "reteica")):
        return "La retención en la fuente y las retenciones territoriales son mecanismos de recaudo anticipado. Para determinar si aplican hay que identificar el concepto, la calidad tributaria de las partes, la base y las reglas vigentes. No asumiré tarifas sin verificar la norma y el año."
    if any(t in q for t in ("renta", "declaración de renta", "declaracion de renta")):
        return "El impuesto sobre la renta grava la renta fiscal determinada conforme a las reglas aplicables al contribuyente. Debemos identificar si es persona natural o jurídica, el año gravable, ingresos, costos, deducciones, rentas exentas, descuentos y soportes. No puedo calcular un impuesto definitivo sin esos datos ni verificar normas actualizadas desde este chat local."
    if "ica" in q or "industria y comercio" in q:
        return "El ICA es un tributo territorial relacionado con actividades industriales, comerciales o de servicios. Su aplicación depende del municipio, la actividad, el período y las reglas locales. Para analizarlo hay que conocer el municipio y el hecho económico."
    if any(t in q for t in ("gmf", "4x1000", "cuatro por mil")):
        return "El GMF se relaciona con determinadas disposiciones de recursos a través del sistema financiero. Su causación y exenciones dependen de los hechos y requisitos legales aplicables."
    if "iva" in q:
        return "El IVA es un impuesto sobre el valor agregado. Para evaluar su causación y si procede como descontable, se deben revisar la operación, la calidad de los responsables, los requisitos legales y los soportes. Un cálculo matemático no determina por sí solo el tratamiento fiscal."
    if any(t in q for t in ("impuesto", "tribut", "fiscal")):
        return "Puedo orientarte en impuestos nacionales y territoriales de Colombia: renta, IVA, retención en la fuente, ICA y GMF. Indícame el impuesto, el año gravable, si se trata de persona natural o empresa y el municipio cuando corresponda. Esta versión no consulta normativa vigente automáticamente."
    if any(t in q for t in ("nómina", "nomina", "pila", "seguridad social", "cesantías", "cesantias", "prestaciones")):
        return "En nómina y seguridad social podemos revisar salarios, novedades, IBC, aportes PILA, prestaciones y soportes. Para una validación concreta necesito el período, tipo de trabajador y datos del caso. Esta demostración no verifica tarifas o normas vigentes en línea."
    if any(t in q for t in ("niif", "estado financiero", "estados financieros", "deterioro", "provisión", "provision")):
        return "En contabilidad bajo NIIF debemos evaluar reconocimiento, medición, presentación y revelación según el marco aplicable a la entidad. Para revisar un caso necesito conocer la operación, los soportes y el grupo contable."
    if "factura" in q or "soporte" in q:
        return "La factura y los demás soportes permiten comprobar la realidad, trazabilidad y tratamiento de una operación. Su validez debe verificarse antes de registrarla definitivamente."
    if "riesgo" in q or "auditor" in q:
        return "En una revisión preliminar, identifica soportes faltantes, posibles errores de clasificación, obligaciones tributarias y controles de aprobación. La conclusión definitiva exige evidencia suficiente."
    if "costo" in q:
        return "El costo representa recursos consumidos o atribuibles a la adquisición o producción, según la naturaleza de la operación. Para clasificarlo correctamente hacen falta detalles y soportes."
    if "excel" in q or "informe" in q or "reporte" in q:
        return "El informe Excel de CONTAVERSE IA contiene tres hojas: Resumen ejecutivo, Análisis multiagente y Validación humana. Se genera a partir de las decisiones guardadas en el historial local."
    return ("Todavía no puedo responder esa pregunta con las reglas de esta demostración. "
            "Prueba con: '¿Qué es el impuesto de renta?', '¿Qué es el ICA?', '¿Qué significa un asiento cuadrado?' "
            "o 'Resume el último caso'. Para respuestas abiertas necesitaremos conectar un modelo de IA.")

st.markdown("### 💬 Conversa con Gabriela")
st.caption("Chat demostrativo sin API: orientación programada sobre impuestos, contabilidad, costos, auditoría y nómina, más el resumen del último caso analizado. No sustituye asesoría profesional.")
if "gabriela_mensajes" not in st.session_state:
    st.session_state["gabriela_mensajes"] = [
        {"rol": "assistant", "texto": "¡Bienvenido! Soy Gabriela. Puedo orientarte en contabilidad, impuestos nacionales y territoriales, costos, auditoría, riesgos y temas laborales. También puedo resumir el último caso analizado."}
    ]
with st.container(border=True):
    for mensaje in st.session_state["gabriela_mensajes"][-12:]:
        with st.chat_message(mensaje["rol"]):
            st.write(mensaje["texto"])
    with st.form("form_gabriela", clear_on_submit=True):
        pregunta_gabriela = st.text_input("Tu pregunta para Gabriela", placeholder="Ej.: ¿Qué es el impuesto de renta en Colombia?")
        enviar_gabriela = st.form_submit_button("Enviar a Gabriela 💬")
    if enviar_gabriela:
        if pregunta_gabriela.strip():
            st.session_state["gabriela_mensajes"].append({"rol": "user", "texto": pregunta_gabriela.strip()})
            respuesta = responder_gabriela(pregunta_gabriela, st.session_state.get("analisis", []))
            st.session_state["gabriela_mensajes"].append({"rol": "assistant", "texto": respuesta})
            st.rerun()
        else:
            st.warning("Escribe una pregunta antes de enviarla.")


# Módulo independiente: diagnóstico tributario orientativo, sin liquidaciones automáticas.
st.markdown("---")
st.header("🏛️ Tributaria IA · Diagnóstico colombiano")
st.caption("Módulo demostrativo con lista de verificación. No consulta normas en tiempo real ni calcula obligaciones definitivas.")
with st.form("form_diagnostico_tributario"):
    a, b, c = st.columns(3)
    with a:
        impuesto = st.selectbox("Impuesto u obligación", ["Impuesto sobre la renta", "IVA", "Retención en la fuente", "ICA", "Impuesto nacional al consumo", "GMF", "Otro / por identificar"])
    with b:
        contribuyente = st.selectbox("Tipo de contribuyente", ["Por definir", "Persona natural", "Persona jurídica"])
    with c:
        anio = st.number_input("Año gravable o período de referencia", min_value=2000, max_value=2100, value=2026, step=1)
    municipio = st.text_input("Municipio (si corresponde)", placeholder="Ej.: Bogotá D. C.")
    actividad = st.text_input("Actividad económica / operación", placeholder="Ej.: comercio al por menor")
    detalle = st.text_area("Describe el caso o la consulta", placeholder="Describe la operación, soportes y dudas. Evita incluir datos personales.")
    ejecutar = st.form_submit_button("🔎 Generar diagnóstico preliminar")

if ejecutar:
    pendientes = []
    if contribuyente == "Por definir":
        pendientes.append("Identificar si el contribuyente es persona natural o jurídica y su régimen aplicable.")
    if not actividad.strip():
        pendientes.append("Precisar la actividad económica y la naturaleza de la operación.")
    if not detalle.strip():
        pendientes.append("Describir la operación y los documentos disponibles.")
    requisitos = {
        "Impuesto sobre la renta": ["Determinar residencia fiscal, calidad del contribuyente y régimen.", "Identificar ingresos, costos, deducciones, beneficios y soportes.", "Verificar reglas y tarifas aplicables al año gravable."],
        "IVA": ["Verificar hecho generador, responsabilidad y tratamiento del bien o servicio.", "Determinar base, tarifa aplicable y requisitos del IVA descontable.", "Revisar facturación y período de declaración."],
        "Retención en la fuente": ["Identificar concepto del pago y calidad del agente retenedor.", "Verificar beneficiario, base mínima, tarifa y excepciones vigentes.", "Comprobar certificados y registros de retención."],
        "ICA": ["Confirmar municipio donde se realiza la actividad gravada.", "Verificar código de actividad, territorialidad, base y tarifa municipal vigente.", "Revisar obligaciones de inscripción, declaración y retención."],
        "Impuesto nacional al consumo": ["Determinar si la operación está sujeta al impuesto.", "Verificar responsables, exclusiones, base y tarifa vigente.", "Comprobar facturación y periodicidad aplicable."],
        "GMF": ["Identificar la transacción financiera y su tratamiento.", "Verificar exenciones y condiciones aplicables.", "Revisar soportes y agente retenedor."],
        "Otro / por identificar": ["Identificar el tributo nacional o territorial pertinente.", "Confirmar hecho generador, sujeto pasivo y autoridad competente.", "Verificar norma, base, tarifa, calendario y soportes."],
    }
    if impuesto == "ICA" and not municipio.strip():
        pendientes.append("Indicar el municipio para revisar la normativa territorial del ICA.")
    st.session_state["diagnostico_tributario"] = {
        "impuesto": impuesto, "contribuyente": contribuyente, "anio": int(anio),
        "municipio": municipio.strip() or "No indicado", "actividad": actividad.strip() or "No indicada",
        "descripcion": detalle.strip() or "No indicada", "verificaciones": requisitos[impuesto],
        "datos_pendientes": pendientes,
        "estado": "Diagnóstico preliminar; sin liquidación ni validación normativa en línea",
    }

if "diagnostico_tributario" in st.session_state:
    dt = st.session_state["diagnostico_tributario"]
    st.markdown("#### Resultado del diagnóstico preliminar")
    st.write(f"**Obligación:** {dt['impuesto']} · **Contribuyente:** {dt['contribuyente']} · **Período:** {dt['anio']}")
    st.write(f"**Municipio:** {dt['municipio']} · **Actividad:** {dt['actividad']}")
    st.markdown("**Verificaciones necesarias antes de concluir:**")
    for verificacion in dt["verificaciones"]:
        st.write("• " + verificacion)
    if dt["datos_pendientes"]:
        st.warning("Información pendiente: " + " ".join(dt["datos_pendientes"]))
    st.info("Este resultado no establece una tarifa ni una obligación definitiva. Requiere revisión de la normativa colombiana vigente y de los soportes.")
    st.download_button("⬇️ Descargar diagnóstico JSON", data=json.dumps(dt, ensure_ascii=False, indent=2), file_name="CONTAVERSE_IA_Diagnostico_Tributario.json", mime="application/json")



# Plantilla de renta con fórmulas y trazabilidad (personas naturales y jurídicas).
st.markdown("### 📑 Renta IA · Papel de trabajo en Excel")
st.write("Plantilla con hojas separadas para persona natural y jurídica, fórmulas de depuración preliminar y matriz de soportes.")
st.info("No es una declaración de renta ni determina el impuesto a pagar. Las cédulas de personas naturales, conciliaciones fiscales, límites, tarifas y tratamientos especiales requieren revisión normativa y profesional.")
st.download_button(
    "📥 Descargar papel de trabajo de renta (Excel)",
    data=base64.b64decode("UEsDBBQAAAAIADaMTF3rdCf7GAEAAOUCAAAPAAAAeGwvd29ya2Jvb2sueG1svdJBTsMwEAXQq1jeUzuJ1YSoaTdsukNwgok9JqaxJ7IdyKU4BRdDFNQWVmzKbjQjfT1p/ma3+JG9YEyOQseLleQMgybjwlPH52xvGr7bbpb2leKhJzqwxY8htUvHh5ynVoikB/SQVjRhWPxoKXrIaUXxSaQpIpg0IGY/ilLKtfDgAv/MO27TaWIBPHZ8H1KOs9aOAibOjqe96XjBWWyd6fhDY7UyFqEvGqmgqPk3KP4FRNY6jXekZ48hf4kijpAdhTS4KXEmfpPuMSYKwALkOcJ4gSpPqFqZomwqVUEjlZL4b6jnOb6/GafhglWdWNY2/Rrquuwro1DLq7MeaaKYf3xOXWhsBb3tVSWlKuH2Chpx7pU4V3b7AVBLAwQUAAAACAA2jExd5azFF0ICAADmEQAADQAAAHhsL3N0eWxlcy54bWzNmM2O2jAQx1/FMpdWakmcDSvEElaAGqmXPXT3UKnqwRAnRBrbkRO2oae+Q9+wT1LFzteuhJZtlxgunhnsX/4ex5Yn89uSA3pkKk+lCDAZuxgxsZVRKpIA74v44xTfLublLC8OwO53jBWo5CDyWRngXVFkM8fJtzvGaT6WGRMlh1gqTot8LFXi5JliNMqrYRwcz3WvHU5TgSui2POQFznayr0oAkx6QWSaz1GAPdfFyCDXMmIBHn0Yjdybb19Y9P2dtt/f/Pn1GyNnMXdaaIWKpejoPm5CejI/0SOFABOix5UzQTkzoTVVkBay4TUjmnZj+reA6xqwlSAVUskmwGH9+180GRhNfG/iL09D14bJcwrQ5nli8pwCVG1Gi4IpEaYAqLYfDhkLsJCCtcS684uDEkUPxJu8elwuIY2MrmR9ZMbOk/FvxJ+s/Olqej7+p2Xohi/orw29UhupIqbatfJwFzSr3tlO21u/JwzgvjoAvsbP9msZ9/aq3qmiNVOA2jSo2jH0PrJ5RI/u+f+KL+PuOa8HkB6AZhkc7vZ8w1SoDyD9t46GUvS9FKDzVhqm/VMleMfmcGYJxL6EY1kg9rNg5124srQQnn0Jx7JA7Gfh3BLqa87pp5s1EXZ2hX8JeRhURE/C1SXkwZaIE5bC1uYcVMQJSzFoHmyJGPyc1P4S0kRw1l18aRNAPxTNHlhpUOZyW8b2755vLnuYA2Aw2eQSZNeFUK8G0jXRsyKrjaOqLA/wXaURnpY6/ZIq1273xWbxF1BLAwQUAAAACAA2jExd+lwBWQMDAADaDQAAEwAAAHhsL3RoZW1lL3RoZW1lMS54bWy9V9tymzAU/BVG7w03c/OEZBLHbh/SaafJD8ggQI0QHkmOnb/vIG4CjOM0duwHS2LP2UXnsMLXt/ucaK+IcVzQEJhXBtAQjYoY0zQEW5F888HtzTWciwzlSKMwRyFYZFB8//0MtH1OKJ/DEGRCbOa6zqMM5ZBfFRtE9zlJCpZDwa8KluoxgztM05zolmG4eg4xBW3eJUE5ooKXCxFhT9EBsvJa/GKWP/yNLwjTXiEJwQ7TuNg9o70AGoFcLAgLgSE/QNNvrvU2ioiJYCVwJT9NYB0Rv1gykKXrNtJYWv7M7BgkgogxcOmX3y6jRMAoQrSWo4JNxzV8qwErqGp4IHvgmfYgQGGwxwyBe2/N+gESVQ1n4xtdBcsHpx8gUdXQGQXcGdZ9YPcDJKoauqOA2fLOs5b9AInKCKYvY7jr+b7bwFtMUpAfB/GB6xreQ4PvYLrSalUCKnqN9ytJcIRk3+Xwb8FWBRWyylBgqom3DUpgVDYoJHjNsPaI00xIHjhH8B1AxI8C9AFnjum7Ao5QHyFt6ToGXd0MuTW5mHwkE0zIk3gj6JFLcbwgOF5hQuRERrWl2GQLwhrCHjBlsBvzOlXKtU3BQ2CAyVzSQTAV1ZrrNU89nJNt/rOI66Y3WzuAcw5Fd8FwFJ9oGeQs5aqGEneyDs+e0NHRDXXYJ+qQd3KyEN/8sJDgqBBdKQ/BVIPlKeHMarvlESQoLgtWJ+iV9SwlDmZTd2R9dmtPKDHPYIyavMaUkqlm67rwDEVWpHj+YSVBMCGk3KpLFFkf2wGh/Zm2K/m95u7+yyw2jIsHyLMKJy+15ytVaALD+QIaq9yZy9Howz1ESYIiMbHSTR+5qLMcvPxZdDkptgKxpyzeaWuyZX9gHALHMx0DaDHmoimAFmPWtc/4/aJbh2STwdrJew9thZfjllMRK+UMpffnteJ1ujrLcfV+1MC1puzWm34SL3A+Bsq5pPhH4H/UUyurPPexqepQ5U0arT0hz76Q0XZd+XWGOmzZ0mOb1zE5G/yBalZu/gFQSwMEFAAAAAgANoxMXQ0euehlAAAAcwAAABQAAAB4bC9zaGFyZWRTdHJpbmdzLnhtbAXBUQrDIAwA0KtI/mfcPsaQ2p5F2rQKJhaTDY+/95ZtcnM/Glq7JHj6AI5k70eVK8HXzscHtnWZUdXc5CYaZ4JidkdE3QtxVt9vksnt7IOzqe/jQr0H5UMLkXHDVwhv5FwFHK5/UEsDBBQAAAAIADaMTF0cnHSKRgMAAJAJAAAYAAAAeGwvd29ya3NoZWV0cy9zaGVldDEueG1slVZLcuM2EL1KF9Yzpv7+ZOQpRZaTVE1slz0zWbfJJgUHQLMAUNZM5Qo5RJZZeOUj6GIpUhJFYxiptCIBvtf90A8N8MPHpVawIOskm7HonnQEkIk5kSYbi8Kn78/Ex8sPy4tntn+6OZGHpVbGXSzHYu59fhFFLp6TRnfCOZmlVilbjd6dsM0il1vCpKJpFfU6nVGkURpRBqxmryvwnYWEUiyUv+fnX0lmcz8W3aGAqATGrNzmCVqWIgVoXFbPZ5n4+Vj0zwXMZZKQGYuOgLhwnvUf62/dXZg1vbeh92p6r38Evb+h93fZB0fQBxv6YJf97Aj6cEMf7ujDI+ijDX10BD3aWVB5doUey4HlZ7AVqLSrLMKaXBtY2RyXmElXgKtK5sfCeVt9WVxOb28+T77O7h9m8NsE/oL72c3nCbyHX76s/p6UeRfr7HWcn+s4UT03bZm7apmbtcxdN+eiakmNlfUbCwil31nOV69Oem7VGcKvKC8sxnL1aiC3pKSWBstdD5aMxzcxftAxaOgYhCXkR7K+sNgqI0TfkXVs0IHBkqPIwTd4KuzqJZExuv0yhg0ZwyDwF9dehxB3JZXMyMSS4OH20y0ohDk/IcRsLbmcTSLJeNovZNQQMgoSzIy3mARL2aoJwVNSCTrA74Ui9xMsULElB2QgJ8cOYlasHyUaPlCa04ai0yDJPblCeUy4XVMIv169Wl2oUlbhWa/+8aUz78AwJBQr3O6ihFJppJeLA5vnrCHtLMg1SRZkfWlH+/YJ8TcMMaq4UAhS5wU5z4AQo80YjATMlYwRPFqZHtpM5w1Z50Gar6hksl5mq6wQP2WTSqvRgqmuE7nAd5CTXb1wwvANHOdsPR1Q1H1zXnXbO2fbOO2HU3dP1zupcyVTGWOClZuWSOcKvyPEq3+TynEjQa1etDystdfU2vsfrXVjt6vt7VGbkSGLClzxRB5Lk9nEUsnN51S6OCjCjxqbx2c3PBCvi7LP21uiG4InmaVsa+/O0LJRq9PjYZ/BUXBtabIZTUmtb7R6BJbSUuhFeSesb4O3yBwz+h1tJo0DRakfi87JqQC7vuyqd8959TYU8Mjes96O5oQJ2XLUF5Ay+3qwzlT/XF3+B1BLAwQUAAAACAA2jExd2OPLdV4EAACwDgAAGAAAAHhsL3dvcmtzaGVldHMvc2hlZXQyLnhtbJVXy27bOBT9lQuuZoA4tuR3ELmw4xQToE0Mx9NZ09KVzVQiVZJynGJ+qItZDOYT8mMDUo4sK7RbryQ+Du/huVdH5PWHbZrABqViggfEu2wRQB6KiPFVQHIdNwbkw+h6e/Us5Fe1RtSwTROurrYBWWudXTWbKlxjStWlyJBv0yQWMqVaXQq5aqpMIo0sLE2afqvVa6aUcWIWtL1fGD6rgxaYQEshvprGXRSQlp2dUY7w8pglTAekT0CL7BPG+gaTJCDjAQEaarbBGeUYkKXQWqRmnIDSVGNAYim+IyfQHF03K9EOW3smH+0mZhIijGme6Ll4/gPZaq0D4nXtKturUCQWEIoEUmbEI5DSrX0+s0ivA9IeElizKEJu9gFhrrRI/yrGvP0yBdzfwf0S7rfPgLd38PY+eucMeGcH7+yjD86Ad3fw7h7ePQPe28F7Z8Cb+xTYnE2ppqYhxTNIO8mky4hQgMsE2oIKzZyxR0BZyXRAlJZ2ZDOa394vxvA3zG7njw/3Y7gfL/6cjz+ZiJsibrnCpFyhWfbdOPqmjr5bR9/Hal/Tbqayp3aFep30gmUCIoRQcC3ZMn9BrtFJuW1D+K3aAjOUSnAKnOpc0uQA+o5Jp8KkU1to/PqvgJWkG7pM3Aw6VQaFHWxGfsvvnQ7arQTt1nP2+mPFUuTQBLFUKDc0ZIKjcsbvHlFASGOFLGYhlaep9Ivq8vunqqtvw/RrUW4EDzHTwsnMDflCEyHh5mHmwty4MY8iE1JjVY/X//jpXQ0q5Ae19e74SqISCpYy10JBzFRIkyMCDywl7zDBLSf7QZkLV8UPK4yGNUZT3IgkL9J8ARKX9IkqeIEIVZgj18LNbfjr3IYnuXmtqpO0imU7x2TjeIzQG7RbYRSPJoPGZGhmxqcY7rC+f4Tigdl5R7kJYxtKM51rthHK+IhErqmbr/frCu7mHpXQr/Lz330pqlppkEkRYmR87YiQ/hnE/NPEqlbr1b12ilEeFoV3Brv2Gezap9lV7deUnKvy5iaBkLz+8y1nEYVMYsJSxmvOVrLrOIvQazUmnteYeH5j4rV/Xo6d0+VYdfBdqBphBbgtnip/QvO0e0jZUVm7Z8i6N34nvV6VXs8t64QqrIhp3CbLX3/Yb2ZH3c2z5xa405h43Z8L2zstbL/KvP4rmJaGCPZooKlkQsFv9w8gUWkqjcYUllTh727u/TM07p/WuOroXt3S7z7PHuaL8f3i1s2jPn+OKk80jUQlI5cwy9Ea2LccJTMZynJZ/P4gxChPqLyAUKQZcrU7JVyUNWZGeMgStgMUnze8gNBSKKBPudKoLt1bH7r37Fd/FO9OHbePi/H0wblfxxmNR8xYTOHQG6YsSW6vK2xD4cXYUIzmNkcTN8taZt4uQG/H5xTlCs29yp6syxZIjE3GrswJtYAfzszoCj9TuWJcQYKxDkjrsk9AFsci+65FZt+6BIor2ltrjTRCaVptArEQumwUkcrL5+h/UEsDBBQAAAAIADaMTF1A4SR3ZQQAALQOAAAYAAAAeGwvd29ya3NoZWV0cy9zaGVldDMueG1slVfLcts2FP2VO1i1M9aD1NtjKqOHM3WmsTWym64h8lKCQwIMAMpyp7/QD8kii04/QT/WAShLFAMp0YrE4+AenHt5CNy826QJrFEqJnhAvHqTAPJQRIwvA5LruNYn74Y3m+sXIT+rFaKGTZpwdb0JyErr7LrRUOEKU6rqIkO+SZNYyJRqVRdy2VCZRBpZWJo0/Gaz20gp48QsaHs/MXxRRy0wgRZCfDaNuyggTTs7oxzh9TFLmA5Ij4AW2e8Y6wkmSUBGfQI01GyNM8oxIAuhtUjNOAGlqcaAxFL8hZxAY3jTKEU7bh2YvLebmEmIMKZ5oufi5Tdky5UOiNexq2yuQ5FYQCgSSJkRj0BKN/b5wiK9CkhrQGDFogi52QeEudIi/bMY8w7LFHB/B/f3cL91Aby1g7cO0dsXwNs7ePsQvX8BvLODdw7wzgXw7g7evQDeOKTA5mxKNTUNKV5A2kkmXUaEArxPoC2o0MwZeQSUlUwHRGlpR9bD+e390wj+htnt/PHhfgQf/phv/5neTUYm5rqIvF9jvF+jse+bOPqmjr5bR9/7cl/Dbqe0q1aJfJX2E8sERAih4FqyRf6KXKOTcsuG8JuVBWYoleAUnnO5/RaxkB6Bv+PSLnFpV5Yabf8VsJR0TReJm0O7zKGwhPXQb/rd80E7paCdat62X5csRQ4NEAuFck1DJjgqZ/zOCQ2ENHbIYhZSeZ5Kr6gwv3euwno2TK8SZSJ4iJkWTmZuyCeaCAmTh5kLM3FjHkUmpMayHtv/+Pld9Uvk+5X17vhSohIKFjLXQkHMVEiTEwL3LSXvOMFNJ/v+Pheumh+UGA0qjKa4FklepPkKJC7oM1XwChGqMEeuhZvb4Oe5Dc5y85plN2kWy7ZPycbxFKE3aKfEKB6O+7XxwMyMzzHcYX3/BMUjw/NOchPGOJRmOtdsLZRxEolcUzdf7+cV3M09KaFf5ud/96WocqVBJkWIkXG2E0L6FxDzzxMrm61XddspRnlYFN4F7FoXsGudZ1e2X1NyrsqbmwRCsv32JWcRhUxiwlLGK862Z9d2FqHXrI09rzb2/NrYa/24HNvny7Hs4LtQFcIKcFM8Vf6M5mn3kLKTsnYukPVg/E563TK9rlvWMVVYEtO4TZZvv9pvZkfdzbPrFrhdG3udHwvbPS9sr8y8+iuY7g0R7OFAU8mEgl/uH0Ci0lQajSksqMJf3dx7F2jcO69x2dG9qqXffZw9zJ9G90+3bh7V+XNUeaJpJEoZqcMsR2tgX3KUzGQoy2Xx+4MQozyh8gpCkWbI1e6UcLWvMTPCQ5awHaD4vOEVhJZCAX3OlUZVd2994N6zX/5RfHfquH18Gk0fnPt1nNJ4xIzFFA69ZsqS5PbKwtYUXo0NxWhudDRxs6xk5u0S9HaETlEu0dyt7Ol63wKJscnYtTmjFvDjmRld4kcql4wrSDDWAWnWewRkcSyy71pk9q1DoLimvbVWSCOUptUiEAuh940i0v4COvwfUEsDBBQAAAAIADaMTF0Sl5PwNwMAAJMLAAAYAAAAeGwvd29ya3NoZWV0cy9zaGVldDQueG1slZbRbtowFIZfxfLVJm0NtoG2qGFiULRK3YrKtEm7c5MTsJbYmW0K3Tvtao/QF5vssGC21BJXcOz/P8f2Z+X46t2uKtEjaCOUTDE562EEMlO5kKsUb2zx9gK/G1/tRlulv5s1gEW7qpRmtEvx2tp6lCQmW0PFzZmqQe6qslC64tacKb1KTK2B595WlQnt9YZJxYXELqEf/SJga44i5Ao9KPXdBTd5inteXXMJ6GlZl8KmmGFkVX0LhZ1CWaZ40seIZ1Y8woJLSPGDslZVbh4jY7mFFBda/QSJUTK+SoJqx9FhJXO/iYVGORR8U9p7tf0AYrW2KSYDn2U3ylTpDZkqUSXc4WFU8Z3/3YrcrlPMLjFaizwH6faBso2xqvrazJFDmsZO93ba2ik7wc72dnao3j/B3t/b+4fqFyfYB3v74GAfnGAf7u3DE+zJAYFnNuOWu0CrLdJe5HC5Q2jMLUB/oTKnmRCMTHOfUmys9jOP48/3k2+T9ze3N7PJDM2u0fJucXf/+XrpSj42hdsU79sUSTs27RibdYxdd4zNw7HE7ybYlLsQtkHz8qaYz3D+z6amSmZQW9W5iRctVouHzRNIC12+abfvCy+VRtO7RZdn1u2ZqWxTgbQKGVUr3V3vutt7DwVokJngKEGbB5HxTDz/ll0Z5i+uWOQ8V+jV8vlX8km9PvL+h6F/OO1p3yckvW5gg0A5iCqHgXIYVZ4HyvOo8iJQXkSVl4HyMqokzQd5f9F7cW1wLaeExLU01NK4loVaFteGsEicFglxuQ99TBsCI3FiJERG4sxICI3EqZEQG4lzoyE3GudGQ240zo2G3GicGw25ue4W04bcaJwbDbm5zhHThtxonBsNudE4Nxpycx/omDbkRuPcWMiNxbmxkBuLc2MhNxbnxkJurJNb8k8XrkCvwD3PzFGENBSu745cn2vsx8qcW958ja1Q0qBMbWTb344nkX2qIcWlMO6d98NnnvdH8/1q3UN0U3IyxsvnX28+qTcLkLlwrQy7su20C44Tdwz5tdV8BR+5XglpUAmFTXHv7Bwj3bRg/9+q2v8bYNS8Qv9Ga+A5aBcxjAqlbBs0p9C+r8d/AFBLAwQUAAAAAAA2jExdgaQgmygBAAAoAQAACwAAAF9yZWxzLy5yZWxz77u/PD94bWwgdmVyc2lvbj0iMS4wIiBlbmNvZGluZz0idXRmLTgiPz48UmVsYXRpb25zaGlwcyB4bWxucz0iaHR0cDovL3NjaGVtYXMub3BlbnhtbGZvcm1hdHMub3JnL3BhY2thZ2UvMjAwNi9yZWxhdGlvbnNoaXBzIj48UmVsYXRpb25zaGlwIFR5cGU9Imh0dHA6Ly9zY2hlbWFzLm9wZW54bWxmb3JtYXRzLm9yZy9vZmZpY2VEb2N1bWVudC8yMDA2L3JlbGF0aW9uc2hpcHMvb2ZmaWNlRG9jdW1lbnQiIFRhcmdldD0iL3hsL3dvcmtib29rLnhtbCIgSWQ9IlI4YTg4MDQ2OTJiMTQ0MzRlIiAvPjwvUmVsYXRpb25zaGlwcz5QSwMEFAAAAAgANoxMXRZotBxDAQAAzwQAABoAAAB4bC9fcmVscy93b3JrYm9vay54bWwucmVsc83UMW7DIBQG4KtY7DXY4JhUcbJ06ZrmAs/wsK0YsAxpnbN16JF6happVdlVhy6RsjD8SL8+HoL317fNbrJ98oxj6LyrSJYykqBTXneuqcgpmjtJdtvNHnuInXeh7YaQTLZ3oSJtjMM9pUG1aCGkfkA32d740UIMqR8bOoA6QoM0Z2xFx3kHWXYmh/OA/2n0xnQKH7w6WXTxj2Ia4rnHQJIDjA3GitCp/87SyfYkedQV2Wsj1notoMhzEFIyktCrgWKLFpeeS/S1ZjNVaUooeCZKVRdCAlxTFVoYUT/FsXPN72nNt2Y8VZc1WxWZ5KYQXMtr8l78eAwtYlzSfuLPAyDG+fSkUUIbhDqTTEBW3gAvn1+u0FkuueAgmRAMb4DHZzxjZL2CssxrrgWqq76If/LEgmc41KYWnDGRw/rCo4tvafsBUEsDBBQAAAAIADaMTF3DxSAoIgEAAO4EAAATAAAAW0NvbnRlbnRfVHlwZXNdLnhtbM2UwUoDMRCGX2XJVZq0VUSk2x7Uqwr6AiE7uxuaTEJmum6fzYOP5CtIUykiwlLcQi+Zy+T/vv8yn+8fi1XvXdFBIhuwFDM5FQWgCZXFphQbric3YrVcvG4jUNF7h1SKljneKkWmBa9JhgjYe1eH5DWTDKlRUZu1bkDNp9NrZQIyIE94lyGWi3uo9cZx8dAz4B7beyeKu/3eDlUKHaOzRrMNqDqsfkEmoa6tgSqYjQdkSTGBrqgFYO9kntJrixc5WP3JTODoOOh3K5nA5R1qbaQD4qmDlGwFxbNO/Kg9lEL1ThFvHZAcuWEOHUJzCx727+zfAjlmsGyrE1QvnCw2o3f+mT0k8hbSOn8klcdsZJlD/rEi83MRuTwXkauTi6h8vZZfUEsBAhQDFAAAAAgANoxMXet0J/sYAQAA5QIAAA8AAAAAAAAAAAAAAKSBAAAAAHhsL3dvcmtib29rLnhtbFBLAQIUAxQAAAAIADaMTF3lrMUXQgIAAOYRAAANAAAAAAAAAAAAAACkgUUBAAB4bC9zdHlsZXMueG1sUEsBAhQDFAAAAAgANoxMXfpcAVkDAwAA2g0AABMAAAAAAAAAAAAAAKSBsgMAAHhsL3RoZW1lL3RoZW1lMS54bWxQSwECFAMUAAAACAA2jExdDR656GUAAABzAAAAFAAAAAAAAAAAAAAApIHmBgAAeGwvc2hhcmVkU3RyaW5ncy54bWxQSwECFAMUAAAACAA2jExdHJx0ikYDAACQCQAAGAAAAAAAAAAAAAAApIF9BwAAeGwvd29ya3NoZWV0cy9zaGVldDEueG1sUEsBAhQDFAAAAAgANoxMXdjjy3VeBAAAsA4AABgAAAAAAAAAAAAAAKSB+QoAAHhsL3dvcmtzaGVldHMvc2hlZXQyLnhtbFBLAQIUAxQAAAAIADaMTF1A4SR3ZQQAALQOAAAYAAAAAAAAAAAAAACkgY0PAAB4bC93b3Jrc2hlZXRzL3NoZWV0My54bWxQSwECFAMUAAAACAA2jExdEpeT8DcDAACTCwAAGAAAAAAAAAAAAAAApIEoFAAAeGwvd29ya3NoZWV0cy9zaGVldDQueG1sUEsBAhQDFAAAAAAANoxMXYGkIJsoAQAAKAEAAAsAAAAAAAAAAAAAAKSBlRcAAF9yZWxzLy5yZWxzUEsBAhQDFAAAAAgANoxMXRZotBxDAQAAzwQAABoAAAAAAAAAAAAAAKSB5hgAAHhsL19yZWxzL3dvcmtib29rLnhtbC5yZWxzUEsBAhQDFAAAAAgANoxMXcPFICgiAQAA7gQAABMAAAAAAAAAAAAAAKSBYRoAAFtDb250ZW50X1R5cGVzXS54bWxQSwUGAAAAAAsACwDVAgAAtBsAAAAA"),
    file_name="CONTAVERSE_IA_Papel_Trabajo_Renta_Natural_Juridica.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    key="descargar_renta_natural_juridica",
)
st.caption("Abre el Excel, selecciona la hoja del contribuyente, diligencia las celdas de entrada y registra los soportes. Los resultados se calculan mediante fórmulas.")

st.header("📁 Cargar caso empresarial")
st.write("Carga un CSV con columnas: descripcion, base, iva_porcentaje, soporte_factura, retencion_definida.")
archivo = st.file_uploader("Selecciona tu archivo CSV", type=["csv"])
simular_fallo = st.checkbox("🧪 Simular fallo del motor primario (prueba Failover Cascade)", value=False, help="Demostración controlada: el motor primario no tiene API conectada. El respaldo usa reglas locales.")
if st.button("🚀 Analizar caso", type="primary", disabled=archivo is None):
    try:
        filas = leer_csv(archivo)
        if not filas:
            st.error("El archivo no contiene operaciones.")
        elif not {"descripcion", "base", "iva_porcentaje"}.issubset(filas[0]):
            st.error("Faltan columnas obligatorias: descripcion, base, iva_porcentaje.")
        else:
            st.session_state["analisis"], st.session_state["traza_failover"] = analizar_con_failover(filas, simular_fallo)
            st.session_state["archivo"] = archivo.name
    except (UnicodeDecodeError, csv.Error, ValueError) as exc:
        st.error(f"No se pudo analizar el archivo: {exc}")

if "analisis" in st.session_state:
    st.success(f"Análisis completado: {len(st.session_state['analisis'])} operación(es).")
    st.markdown("### 🔁 Evidencia de Failover Cascade")
    st.dataframe(st.session_state.get("traza_failover", []), hide_index=True, use_container_width=True)
    st.caption("Prueba de continuidad con fallo inducido y respaldo local. No acredita conexión real a un modelo externo ni comunicación con Frame.")
    st.download_button("⬇️ Descargar evidencia Failover (JSON)", data=json.dumps({"prueba": "simulada", "eventos": st.session_state.get("traza_failover", [])}, ensure_ascii=False, indent=2), file_name="CONTAVERSE_IA_Evidencia_Failover.json", mime="application/json")
    for caso in st.session_state["analisis"]:
        st.subheader(f"Caso {caso['fila']}: {caso['descripcion']}")
        cols = st.columns(3)
        for col, clave, titulo_agente in zip(cols, ["contable", "tributario", "riesgo"], ["🧾 Contable IA", "🏛️ Tributario IA", "🛡️ Riesgos y Auditoría IA"]):
            with col:
                with st.container(border=True):
                    st.subheader(titulo_agente)
                    st.markdown("**Resultado del análisis**")
                    st.write(caso[clave]["conclusion"])
                    estado = caso[clave]["estado"]
                    if estado == "Alerta documental":
                        st.error(f"🔴 {estado}")
                    elif estado in ("Cálculo preliminar", "Por validar"):
                        st.info(f"🔵 {estado}")
                    else:
                        st.caption(estado)

        # Un solo asiento contable preliminar por caso.
        asiento = caso["contable"].get("asiento_preliminar")
        if asiento:
            st.markdown("### 📒 Asiento contable preliminar")
            st.dataframe([
                {"Concepto": "Compra", "Débito": asiento["debito_compra"], "Crédito": 0},
                {"Concepto": "IVA descontable (por validar)", "Débito": asiento["debito_iva"], "Crédito": 0},
                {"Concepto": "Proveedores", "Débito": 0, "Crédito": asiento["credito_proveedor"]},
            ], hide_index=True, use_container_width=True)
            st.caption("Propuesta preliminar para compra a crédito, sin retenciones y suponiendo IVA descontable. Requiere validación contable y tributaria.")
        for alerta in caso["comparacion"]["discrepancias_y_pendientes"]:
            st.warning(alerta)
        st.write("**Dictamen:**", caso["comparacion"]["dictamen"])
        with st.expander("Ver JSON estructurado intercambiado entre agentes"):
            st.json(caso)
        st.caption(f"Modo de ejecución: {caso['modo']}")
        st.markdown("**👤 Validación humana**")
        decision = st.radio("Decisión", ["Aceptar", "Modificar", "Rechazar"], horizontal=True, key=f"decision_{caso['fila']}")
        justificacion = st.text_area("Justificación (obligatoria para Modificar y Rechazar)", key=f"just_{caso['fila']}")
        if st.button("Guardar decisión", key=f"guardar_{caso['fila']}"):
            if decision in ("Modificar", "Rechazar") and not justificacion.strip():
                st.error("Debes escribir una justificación.")
            else:
                registro = {"fecha": datetime.now().isoformat(timespec="seconds"), "archivo": st.session_state.get("archivo", ""), "caso": caso, "decision": decision, "justificacion": justificacion.strip()}
                try:
                    with LOG.open("a", encoding="utf-8") as f:
                        f.write(json.dumps(registro, ensure_ascii=False) + "\n")
                    st.success("Decisión guardada en decisiones.jsonl")
                except OSError as exc:
                    st.error(f"No se pudo guardar la decisión: {exc}")
        st.divider()

st.header("📋 Registro de decisiones")
if LOG.exists():
    try:
        registros = leer_registros()
        st.download_button("📊 Descargar informe empresarial Excel", data=generar_excel(registros), file_name="CONTAVERSE_IA_Informe_Empresarial.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except (OSError, ValueError) as exc:
        st.error(f"No se pudo preparar el registro de decisiones: {exc}")
else:
    st.caption("El registro se creará cuando guardes la primera decisión.")




def generar_excel_analytics_avance5(resultados):
    """Excel ejecutivo de los resultados realmente obtenidos en esta sesión."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen ejecutivo"
    detalle = libro.create_sheet("Detalle 60 casos")
    metodo = libro.create_sheet("Metodología")
    azul, azul_claro, dorado = "162B53", "EAF1FC", "E8B65A"

    def cabecera(hoja, texto, ancho):
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
        c = hoja.cell(1, 1, texto)
        c.fill = PatternFill("solid", fgColor=azul)
        c.font = Font(name="Aptos Display", size=17, bold=True, color="FFFFFF")
        c.alignment = Alignment(vertical="center", indent=1)
        hoja.row_dimensions[1].height = 40
        hoja.sheet_view.showGridLines = False

    def encabezados(hoja, fila, titulos):
        for j, titulo in enumerate(titulos, 1):
            c = hoja.cell(fila, j, titulo)
            c.fill = PatternFill("solid", fgColor=azul)
            c.font = Font(name="Aptos", bold=True, color="FFFFFF")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.row_dimensions[fila].height = 29

    def renglon(hoja, fila, datos):
        for j, dato in enumerate(datos, 1):
            c = hoja.cell(fila, j, dato)
            c.font = Font(name="Aptos", size=10, color="24324A")
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if fila % 2 == 0:
                c.fill = PatternFill("solid", fgColor=azul_claro)
        hoja.row_dimensions[fila].height = 28

    cuenta = Counter(r["resultado"] for r in resultados)
    tp, tn, fp, fn = (cuenta[k] for k in ("TP", "TN", "FP", "FN"))
    dividir = lambda a, b: a / b if b else 0
    precision = dividir(tp, tp + fp)
    recall = dividir(tp, tp + fn)
    f1 = dividir(2 * precision * recall, precision + recall)
    exactitud = dividir(tp + tn, len(resultados))
    promedio = sum(float(r["latencia_motor_ms"]) for r in resultados) / len(resultados)

    cabecera(resumen, "CONTAVERSE IA  |  ANALYTICS 5.0", 5)
    resumen.merge_cells("A2:E2")
    resumen["A2"] = "Informe ejecutivo de prueba masiva • Datos calculados en la ejecución actual"
    resumen["A2"].font = Font(name="Aptos", italic=True, color="52637F", size=11)
    encabezados(resumen, 4, ["INDICADOR", "VALOR", "INDICADOR", "VALOR", "OBSERVACIÓN"])
    metricas = [
        ("Casos analizados", len(resultados), "Exactitud", exactitud, "Casos clasificados correctamente"),
        ("Precisión", precision, "Recall", recall, "Métricas de detección de alertas"),
        ("F1", f1, "Latencia media (ms)", round(promedio, 4), "Solo motor local, sin interfaz"),
        ("Verdaderos positivos (TP)", tp, "Verdaderos negativos (TN)", tn, "Comparación con Ground Truth"),
        ("Falsos positivos (FP)", fp, "Falsos negativos (FN)", fn, "FN: alertas esperadas no detectadas"),
    ]
    for i, datos in enumerate(metricas, 5):
        renglon(resumen, i, datos)
    for celda in ("D5", "B6", "D6", "B7"):
        resumen[celda].number_format = "0.00%"
    resumen["A12"] = "LECTURA DEL RESULTADO"
    resumen["A12"].font = Font(bold=True, color=azul, size=12)
    resumen.merge_cells("A13:E14")
    resumen["A13"] = (f"Se clasificaron correctamente {tp+tn} de {len(resultados)} casos. "
                       f"Se registraron {fn} falsos negativos y {fp} falsos positivos. "
                       "Las anomalías de importes atípicos son criterios pedagógicos de este conjunto de prueba.")
    resumen["A13"].alignment = Alignment(wrap_text=True, vertical="center")
    resumen["A13"].fill = PatternFill("solid", fgColor="FFF5E0")
    resumen.row_dimensions[13].height = 36
    for col, ancho in {"A":30,"B":17,"C":29,"D":18,"E":46}.items():
        resumen.column_dimensions[col].width = ancho
    resumen.freeze_panes = "A5"

    cabecera(detalle, "CONTAVERSE IA  |  DETALLE DE CASOS", 7)
    detalle.merge_cells("A2:G2")
    detalle["A2"] = "TP: verdadero positivo  ·  TN: verdadero negativo  ·  FP: falso positivo  ·  FN: falso negativo"
    detalle["A2"].font = Font(italic=True, color="52637F")
    columnas = ["ID caso", "Categoría", "Alerta esperada", "Alerta detectada", "Clasificación", "Latencia (ms)", "Detalle técnico"]
    encabezados(detalle, 4, columnas)
    for i, r in enumerate(resultados, 5):
        renglon(detalle, i, [r["id_caso"], r["categoria_caso"], r["gt_alerta"], r["pred_alerta"], r["resultado"], float(r["latencia_motor_ms"]), r["detalle"]])
        detalle.cell(i, 6).number_format = "0.0000"
        if r["resultado"] in ("FN", "FP"):
            detalle.cell(i, 5).fill = PatternFill("solid", fgColor="FCE8E6")
            detalle.cell(i, 5).font = Font(bold=True, color="A61B1B")
        elif r["resultado"] in ("TP", "TN"):
            detalle.cell(i, 5).font = Font(bold=True, color="147A55")
    for col, ancho in {"A":14,"B":26,"C":21,"D":21,"E":18,"F":19,"G":76}.items():
        detalle.column_dimensions[col].width = ancho
    detalle.freeze_panes = "C5"
    detalle.auto_filter.ref = f"A4:G{4+len(resultados)}"

    cabecera(metodo, "CONTAVERSE IA  |  METODOLOGÍA", 3)
    encabezados(metodo, 4, ["ASPECTO", "DESCRIPCIÓN", "ALCANCE"])
    notas = [
        ("Fuente", "Dataset sintético con Ground Truth declarado", "No contiene datos reales de contribuyentes"),
        ("Motor", "Motor determinista local de CONTAVERSE IA", "Sin API externa en este lote"),
        ("Exactitud", "(TP + TN) / total", "Proporción de aciertos"),
        ("Precisión", "TP / (TP + FP)", "Fiabilidad de alertas emitidas"),
        ("Recall", "TP / (TP + FN)", "Cobertura de alertas esperadas"),
        ("F1", "2 × precisión × recall / (precisión + recall)", "Balance de precisión y recall"),
        ("Latencia", "Tiempo de ejecución del motor por caso", "No mide red, carga, interfaz o revisión humana"),
        ("Importes atípicos", "Alerta pedagógica para valores ≥ $1.000 millones", "No equivale a una obligación tributaria"),
        ("Control humano", "Aceptar, Modificar o Rechazar en el flujo individual", "El lote no simula decisiones humanas"),
    ]
    for i, fila in enumerate(notas, 5):
        renglon(metodo, i, fila)
        metodo.row_dimensions[i].height = 37
    for col, ancho in {"A":24,"B":62,"C":54}.items():
        metodo.column_dimensions[col].width = ancho
    metodo.freeze_panes = "A5"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()

# ===== AVANCE 5: MÓDULO AISLADO DE PRUEBAS MASIVAS =====
# Este bloque no altera el flujo de análisis individual ni sus registros.
import time as _av5_time

st.divider()
st.header("📈 Avance 5 · Prueba de estrés y Analytics")
st.info("Carga aquí el dataset de 60 casos con Ground Truth. Este módulo ejecuta cada fila por separado usando el motor local existente; no envía datos a una API externa ni guarda decisiones humanas automáticamente.")
_av5_archivo = st.file_uploader("Cargar dataset de casos con Ground Truth (CSV)", type=["csv"], key="av5_dataset")
if _av5_archivo is not None:
    try:
        _av5_texto = _av5_archivo.getvalue().decode("utf-8-sig")
        _av5_filas = list(csv.DictReader(io.StringIO(_av5_texto)))
        _av5_campos = {"id_caso", "categoria_caso", "descripcion", "base", "iva_porcentaje", "soporte_factura", "retencion_definida", "gt_alerta"}
        _av5_faltantes = _av5_campos.difference(set(_av5_filas[0].keys()) if _av5_filas else set())
        if _av5_faltantes:
            st.error("Faltan columnas: " + ", ".join(sorted(_av5_faltantes)))
        else:
            st.success(f"Dataset recibido: {len(_av5_filas)} casos.")
            if st.button("🚀 Ejecutar prueba masiva", key="av5_ejecutar", type="primary"):
                _av5_resultados = []
                for _av5_f in _av5_filas:
                    _av5_inicio = _av5_time.perf_counter()
                    _av5_error = ""
                    try:
                        _av5_caso = analizar([_av5_f])[0]
                        _av5_alerta = int(bool(_av5_caso["comparacion"]["discrepancias_y_pendientes"]))
                        _av5_detalle = "; ".join(_av5_caso["comparacion"]["discrepancias_y_pendientes"]) or "Sin alerta del motor"
                    except (ValueError, TypeError, KeyError) as _av5_exc:
                        _av5_alerta = 1
                        _av5_error = str(_av5_exc)
                        _av5_detalle = "Error de validación: " + _av5_error
                    _av5_ms = round((_av5_time.perf_counter() - _av5_inicio) * 1000, 4)
                    _av5_gt = str(_av5_f.get("gt_alerta", "")).strip()
                    if _av5_gt not in ("0", "1"):
                        st.error(f"Ground Truth inválido en {_av5_f.get('id_caso')}: debe ser 0 o 1.")
                        st.stop()
                    _av5_gt_num = int(_av5_gt)
                    _av5_resultados.append({
                        "id_caso": _av5_f.get("id_caso", ""),
                        "categoria_caso": _av5_f.get("categoria_caso", ""),
                        "gt_alerta": _av5_gt_num,
                        "pred_alerta": _av5_alerta,
                        "resultado": "TP" if _av5_gt_num == 1 and _av5_alerta == 1 else "TN" if _av5_gt_num == 0 and _av5_alerta == 0 else "FP" if _av5_gt_num == 0 else "FN",
                        "latencia_motor_ms": _av5_ms,
                        "detalle": _av5_detalle,
                    })
                st.session_state["av5_resultados"] = _av5_resultados
                st.session_state["av5_nombre"] = _av5_archivo.name
    except UnicodeDecodeError:
        st.error("El archivo debe estar en CSV con codificación UTF-8.")

if st.session_state.get("av5_resultados"):
    _av5_r = st.session_state["av5_resultados"]
    _av5_conteo = Counter(r["resultado"] for r in _av5_r)
    _av5_tp, _av5_tn = _av5_conteo["TP"], _av5_conteo["TN"]
    _av5_fp, _av5_fn = _av5_conteo["FP"], _av5_conteo["FN"]
    _av5_div = lambda a, b: a / b if b else 0.0
    _av5_precision = _av5_div(_av5_tp, _av5_tp + _av5_fp)
    _av5_recall = _av5_div(_av5_tp, _av5_tp + _av5_fn)
    _av5_f1 = _av5_div(2 * _av5_precision * _av5_recall, _av5_precision + _av5_recall)
    _av5_accuracy = _av5_div(_av5_tp + _av5_tn, len(_av5_r))
    st.subheader("Resultados de la ejecución real del motor local")
    _av5_cols = st.columns(4)
    _av5_cols[0].metric("Casos", len(_av5_r))
    _av5_cols[1].metric("Exactitud", f"{_av5_accuracy:.2%}")
    _av5_cols[2].metric("Precisión", f"{_av5_precision:.2%}")
    _av5_cols[3].metric("Recall", f"{_av5_recall:.2%}")
    _av5_cols2 = st.columns(5)
    for _av5_c, _av5_et, _av5_val in zip(_av5_cols2, ["TP", "TN", "FP", "FN", "F1"], [_av5_tp, _av5_tn, _av5_fp, _av5_fn, f"{_av5_f1:.2%}"]):
        _av5_c.metric(_av5_et, _av5_val)
    st.caption("Latencia media del motor local por caso: " + f"{sum(r['latencia_motor_ms'] for r in _av5_r)/len(_av5_r):.4f} ms. No incluye interfaz, carga de archivos, red ni validación humana.")
    st.caption("Ground Truth sintético: la regla de alerta para importes atípicos (≥ $1.000 millones) es un criterio pedagógico de prueba, no una obligación normativa. Las anomalías no detectadas aparecen como FN.")
    st.dataframe(_av5_r, use_container_width=True, hide_index=True)
    _av5_salida = io.StringIO()
    _av5_writer = csv.DictWriter(_av5_salida, fieldnames=list(_av5_r[0].keys()))
    _av5_writer.writeheader()
    _av5_writer.writerows(_av5_r)
    st.download_button("⬇️ Descargar resultados reales de la prueba (CSV)", data="\ufeff" + _av5_salida.getvalue(), file_name="CONTAVERSE_IA_Avance5_Resultados_Reales.csv", mime="text/csv", key="av5_descargar")
    st.download_button("📗 Descargar informe profesional Excel (3 hojas)", data=generar_excel_analytics_avance5(_av5_r), file_name="CONTAVERSE_IA_Avance5_Informe_Profesional.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="av5_descargar_excel")
    st.warning("La prueba de estrés automatiza el diagnóstico técnico. La decisión humana (Aceptar, Modificar o Rechazar) sigue disponible en el flujo individual y no se simula como una aprobación real en este lote.")

# ============================================================================
# AVANCE 6 — N11 GEMELO DIGITAL + N12 SIMULADOR DE DILEMAS EMPRESARIALES
# Extensión independiente. No modifica los módulos anteriores.
# Cálculos determinísticos pedagógicos; NO conecta con modelos IA externos.
# ============================================================================
import hashlib as _a6_hashlib

_A6_AREAS = ("Contabilidad", "Tributaria", "Riesgos y Auditoría", "Gobernanza")
_A6_ESCENARIOS = (
    ("Base", 1, 0, 0, 0),
    ("Alternativo 1", 2, 15, 1, 0),
    ("Alternativo 2", 3, 40, 2, 1),
    ("Estrés", 4, 80, 3, 1),
)


def _a6_indicadores(datos, dificultad, presion, contradicciones):
    """Modelo reproducible del estado organizacional, sin aleatoriedad."""
    n = len(datos)
    incompletos = sum(not str(d.get("soporte_factura", "")).strip().lower() in ("si", "sí") for d in datos)
    ret_pend = sum(str(d.get("retencion_definida", "")).strip().lower() not in ("si", "sí", "no") for d in datos)
    valores = [numero(d.get("base", 0)) for d in datos]
    if any(x < 0 for x in valores):
        raise ValueError("Las bases de operación deben ser no negativas.")
    facturacion = round(sum(valores), 2)
    complejidad = dificultad * n + incompletos * 2 + ret_pend + contradicciones * n
    carga = complejidad * (1 + presion / 100)
    capacidad = max(1, n * 3)
    saturacion = round(100 * carga / capacidad, 2)
    riesgo = min(100, round(8 + 9*dificultad + 0.25*presion + 8*contradicciones + (35*incompletos/max(n, 1)) + (15*ret_pend/max(n, 1)), 2))
    tiempo_estimado = round(carga * 4, 2)
    casos_escalados = min(n, int((n * riesgo / 100) + 0.9999))
    if riesgo >= 75 or saturacion >= 120:
        dictamen = "ESCALAR Y APLAZAR decisión hasta revisión de evidencias"
    elif riesgo >= 45 or saturacion >= 90:
        dictamen = "MODIFICAR condiciones y solicitar revisión humana"
    else:
        dictamen = "REVISIÓN HUMANA previa a cualquier aprobación"
    return {
        "Casos": n, "Valor base total": facturacion,
        "Soportes pendientes": incompletos, "Retenciones sin definir": ret_pend,
        "Dificultad": dificultad, "Presión de tiempo (%)": presion,
        "Contradicciones por caso": contradicciones,
        "Carga estimada (unidades)": round(carga, 2),
        "Capacidad de referencia (unidades)": capacidad,
        "Saturación (%)": saturacion,
        "Riesgo modelado (0-100)": riesgo,
        "Tiempo estimado de revisión (min)": tiempo_estimado,
        "Casos sugeridos para escalamiento": casos_escalados,
        "Recomendación automatizada": dictamen,
    }


def _a6_fuente():
    """Acepta dataset N10/Avance5, casos analizados o dataset de ejemplo."""
    if st.session_state.get("a6_csv"):
        return st.session_state["a6_csv"], "CSV cargado para Avance 6"
    if st.session_state.get("analisis"):
        filas = []
        for c in st.session_state["analisis"]:
            filas.append({"descripcion": c.get("descripcion", "Caso"),
                "base": c.get("contable", {}).get("base", 0),
                "soporte_factura": "no" if c.get("riesgo", {}).get("estado") == "Alerta documental" else "si",
                "retencion_definida": "no"})
        return filas, "Casos analizados en sesión (retención representada de forma simplificada)"
    return [
        {"descripcion": "Compra de inventario", "base": "12000000", "soporte_factura": "si", "retencion_definida": "si"},
        {"descripcion": "Servicio sin factura", "base": "7200000", "soporte_factura": "no", "retencion_definida": ""},
        {"descripcion": "Compra de maquinaria", "base": "38000000", "soporte_factura": "si", "retencion_definida": "no"},
        {"descripcion": "Consultoría", "base": "9500000", "soporte_factura": "si", "retencion_definida": "si"},
    ], "Datos sintéticos de ejemplo (no corresponden a operaciones reales)"


def _a6_exportar_csv(filas):
    buf = io.StringIO()
    campos = list(filas[0])
    w = csv.DictWriter(buf, fieldnames=campos, extrasaction="ignore")
    w.writeheader()
    w.writerows(filas)
    return "\ufeff" + buf.getvalue()


st.divider()
st.header("🌐 Avance 6 · Laboratorio Empresarial CONTAVERSE IA")
st.caption("N11: gemelo organizacional lógico · N12: simulador reproducible de dilemas · N13: integración con Frame pendiente.")
st.info("Modelo educativo basado en reglas, datos sintéticos o CSV. Las métricas de riesgo, carga, capacidad y tiempo son índices diseñados para el prototipo; NO son mediciones empíricas, dictámenes legales ni respuestas de una API de IA.")

_a6_upload = st.file_uploader("📁 Dataset organizacional (CSV con descripcion, base, soporte_factura, retencion_definida)", type=["csv"], key="a6_archivo")
if _a6_upload is not None:
    _a6_firma = _a6_hashlib.sha256(_a6_upload.getvalue()).hexdigest()
    if st.session_state.get("a6_archivo_firma") != _a6_firma:
        try:
            _a6_importadas = leer_csv(_a6_upload)
            _a6_obligatorias = {"descripcion", "base", "soporte_factura", "retencion_definida"}
            if not _a6_importadas or not _a6_obligatorias.issubset(set(_a6_importadas[0])):
                raise ValueError("El CSV debe tener datos y columnas: " + ", ".join(sorted(_a6_obligatorias)))
            for _a6_fila in _a6_importadas:
                if numero(_a6_fila["base"]) < 0:
                    raise ValueError("Se encontraron bases negativas.")
            st.session_state["a6_csv"] = _a6_importadas
            st.session_state["a6_archivo_firma"] = _a6_firma
            st.success(f"Se validaron {len(_a6_importadas)} filas para el laboratorio.")
        except (ValueError, UnicodeDecodeError, TypeError) as _a6_error:
            st.session_state.pop("a6_csv", None)
            st.session_state.pop("a6_archivo_firma", None)
            st.error(f"No se pudo cargar el dataset: {_a6_error}")
if st.button("Restablecer fuente sintética / casos de la sesión", key="a6_limpiar"):
    st.session_state.pop("a6_csv", None)
    st.session_state.pop("a6_archivo_firma", None)
    st.rerun()

_a6_datos, _a6_origen = _a6_fuente()
st.caption("Fuente utilizada: " + _a6_origen)
st.markdown("#### N11 · Gemelo Digital Organizacional")
st.write("Áreas: " + " → ".join(_A6_AREAS) + " → revisión y decisión humana.")
_a6_left, _a6_mid, _a6_right = st.columns(3)
with _a6_left:
    _a6_dificultad = st.slider("Dificultad del caso (1–4)", 1, 4, 1, key="a6_dificultad")
with _a6_mid:
    _a6_presion = st.slider("Presión de tiempo (%)", 0, 100, 0, 5, key="a6_presion")
with _a6_right:
    _a6_contradiccion = st.slider("Contradicciones por caso (0–3)", 0, 3, 0, key="a6_contradiccion")
try:
    _a6_estado = _a6_indicadores(_a6_datos, _a6_dificultad, _a6_presion, _a6_contradiccion)
    _a6_kpis = st.columns(4)
    for _a6_col, _a6_k, _a6_formato in zip(_a6_kpis,
        ["Casos", "Saturación (%)", "Riesgo modelado (0-100)", "Casos sugeridos para escalamiento"],
        ["{:.0f}", "{:.1f}%", "{:.1f}/100", "{:.0f}"]):
        _a6_col.metric(_a6_k, _a6_formato.format(_a6_estado[_a6_k]))
    st.progress(min(1., _a6_estado["Riesgo modelado (0-100)"] / 100.))
    st.write("**Interpretación automática por reglas:** " + _a6_estado["Recomendación automatizada"])
    # Una barra por indicador, sin apilar ni sumar capacidad con carga.
    _a6_comparacion = [
        {"Indicador": "Carga estimada", "Unidades": _a6_estado["Carga estimada (unidades)"]},
        {"Indicador": "Capacidad de referencia", "Unidades": _a6_estado["Capacidad de referencia (unidades)"]},
    ]
    st.bar_chart(_a6_comparacion, x="Indicador", y="Unidades", horizontal=True)
    st.caption("Comparación de dos magnitudes independientes del modelo; unidades sintéticas de referencia.")
    with st.expander("🧬 Mapa de variables, relaciones y trazabilidad"):
        st.dataframe([
            {"Variable": "Casos", "Tipo": "Entrada", "Fuente": _a6_origen, "Relación": "Determina carga y capacidad"},
            {"Variable": "Soportes pendientes", "Tipo": "Entrada derivada", "Fuente": "CSV / casos", "Relación": "Aumenta complejidad y riesgo"},
            {"Variable": "Dificultad", "Tipo": "Entrada manipulable", "Fuente": "Control N11", "Relación": "Aumenta carga y riesgo"},
            {"Variable": "Presión de tiempo", "Tipo": "Entrada manipulable", "Fuente": "Control N11", "Relación": "Aumenta carga y riesgo"},
            {"Variable": "Contradicciones", "Tipo": "Entrada manipulable", "Fuente": "Control N11", "Relación": "Aumenta carga y riesgo"},
            {"Variable": "Riesgo y saturación", "Tipo": "Indicador calculado", "Fuente": "Motor local N11", "Relación": "Determina recomendación y escalamiento"},
        ], hide_index=True, use_container_width=True, row_height=62,
            column_config={
                "Variable": st.column_config.TextColumn("Variable", width="medium"),
                "Tipo": st.column_config.TextColumn("Tipo", width="medium"),
                "Fuente": st.column_config.TextColumn("Fuente de datos", width="large"),
                "Relación": st.column_config.TextColumn("Relación entre variables", width="large"),
            })
        st.json(_a6_estado)
    st.markdown("#### N12 · Simulador de Dilemas Empresariales")
    st.write("El laboratorio calcula cuatro condiciones distintas sobre el **mismo conjunto de casos**, sin alterar ni registrar una operación real.")
    _a6_tabla = []
    for _a6_nombre, _a6_d, _a6_p, _a6_c, _a6_crisis in _A6_ESCENARIOS:
        _a6_res = _a6_indicadores(_a6_datos, _a6_d, _a6_p, _a6_c)
        _a6_tabla.append({"Escenario": _a6_nombre,
            "Dificultad": _a6_d, "Presión (%)": _a6_p,
            "Contradicciones": _a6_c, "Riesgo (0-100)": _a6_res["Riesgo modelado (0-100)"],
            "Saturación (%)": _a6_res["Saturación (%)"],
            "Revisión (min)": _a6_res["Tiempo estimado de revisión (min)"],
            "Casos escalados": _a6_res["Casos sugeridos para escalamiento"],
            "Recomendación motor": _a6_res["Recomendación automatizada"]})
    st.dataframe(_a6_tabla, hide_index=True, use_container_width=True,
        row_height=72,
        column_config={
            "Escenario": st.column_config.TextColumn("Escenario", width="medium"),
            "Recomendación motor": st.column_config.TextColumn("Recomendación calculada", width="large"),
        })
    with st.expander("📋 Ver recomendaciones completas por escenario"):
        for _a6_fila in _a6_tabla:
            st.markdown(f"**{_a6_fila['Escenario']}** — {_a6_fila['Recomendación motor']}")
    # Preservar el orden base → alternativos → estrés en el eje horizontal.
    import altair as _a6_alt
    _a6_orden = ["Base", "Alternativo 1", "Alternativo 2", "Estrés"]
    _a6_serie = [{"Escenario": r["Escenario"], "Riesgo modelado": r["Riesgo (0-100)"]} for r in _a6_tabla]
    _a6_grafico = (_a6_alt.Chart(_a6_alt.Data(values=_a6_serie))
        .mark_line(point=True, strokeWidth=3)
        .encode(
            x=_a6_alt.X("Escenario:N", sort=_a6_orden, title="Escenario"),
            y=_a6_alt.Y("Riesgo modelado:Q", title="Riesgo modelado (0–100)", scale=_a6_alt.Scale(domain=[0, 100])),
            tooltip=["Escenario:N", _a6_alt.Tooltip("Riesgo modelado:Q", format=".2f")],
        ).properties(height=320))
    st.altair_chart(_a6_grafico, use_container_width=True)
    st.download_button("⬇️ Descargar matriz reproducible (CSV)", _a6_exportar_csv(_a6_tabla),
        file_name="CONTAVERSE_IA_Avance6_Matriz_Simulacion.csv", mime="text/csv", key="a6_descargar")
    _a6_caso = st.selectbox("Situación empresarial a deliberar", [d.get("descripcion", "Caso") for d in _a6_datos], key="a6_caso")
    _a6_seleccion = st.selectbox("Escenario para la decisión", [r["Escenario"] for r in _a6_tabla], key="a6_escenario")
    _a6_elegido = next(r for r in _a6_tabla if r["Escenario"] == _a6_seleccion)
    st.write("**Contable IA (reglas):** revisar valor y soportes antes de reconocer la operación.")
    st.write("**Tributario IA (reglas):** validar IVA y retenciones, sin asumir tratamiento definitivo.")
    st.write("**Riesgos IA (reglas):** " + _a6_elegido["Recomendación motor"])
    _a6_decision = st.radio("¿Qué decide el estudiante?", ["Aceptar bajo revisión", "Modificar", "Escalar / aplazar"], horizontal=True, key="a6_decision")
    _a6_justificacion = st.text_area("Justificación del estudiante", key="a6_justificacion", placeholder="Explica tu decisión frente a los riesgos detectados")
    if st.button("👤 Registrar deliberación humana N12", key="a6_registrar"):
        if not _a6_justificacion.strip():
            st.warning("Escribe una justificación antes de registrar la deliberación.")
        else:
            _a6_log = BASE / "simulaciones_avance6.jsonl"
            _a6_evento = {"fecha": datetime.now().isoformat(timespec="seconds"),
                "tipo": "SIMULACIÓN EDUCATIVA", "caso": _a6_caso,
                "origen": _a6_origen, "escenario": _a6_elegido,
                "decision_estudiante": _a6_decision, "justificacion": _a6_justificacion.strip(),
                "motor": "Reglas locales; sin IA externa", "sin_efecto_real": True}
            try:
                with _a6_log.open("a", encoding="utf-8") as _a6_f:
                    _a6_f.write(json.dumps(_a6_evento, ensure_ascii=False) + "\n")
                st.success("Deliberación registrada en simulaciones_avance6.jsonl (separado de decisiones reales).")
                st.json({"Decisión estudiante": _a6_decision, "Recomendación motor": _a6_elegido["Recomendación motor"]})
            except OSError as _a6_exc:
                st.error(f"No fue posible registrar: {_a6_exc}")
    with st.expander("📐 Método de cálculo y límites del modelo"):
        st.write("Carga = (dificultad × casos + 2 × soportes pendientes + retenciones sin definir + contradicciones × casos) × (1 + presión / 100).")
        st.write("Capacidad de referencia = máximo(1, 3 × casos). Saturación = 100 × carga / capacidad.")
        st.write("Riesgo = mínimo(100, 8 + 9 × dificultad + 0,25 × presión + 8 × contradicciones + 35 × proporción sin soporte + 15 × proporción sin retención definida).")
        st.write("Tiempo de revisión = 4 minutos × carga. Estos parámetros son **supuestos sintéticos** y deben calibrarse antes de usarlos fuera del aula.")
        st.write("El resultado es reproducible si se conserva el CSV, los parámetros y la versión del código. No es Monte Carlo ni una predicción estadística.")
except (ValueError, TypeError, KeyError) as _a6_exc:
    st.error(f"No se pudo construir el modelo: {_a6_exc}")
st.caption("N13 — Frame: aún no integrado. El gemelo, el simulador y las reglas no tienen actualmente conexión con un mundo 3D ni con un modelo generativo externo.")


# ============================================================================
# N13 · CONSULTA REAL POR HTTP A UNA API REST INDEPENDIENTE
# Agregado sin alterar N11, N12, Avance 5, ni Failover Cascade previos.
# ============================================================================
import urllib.request as _n13_http
import urllib.error as _n13_error


def _n13_config(nombre):
    try:
        valor=st.secrets.get(nombre, "")
    except Exception:
        valor=""
    return str(valor or "").strip()


def _n13_solicitud(url, metodo="GET", contenido=None):
    encabezados={"Accept":"application/json"}
    clave=_n13_config("CONTAVERSE_API_KEY")
    if clave:
        encabezados["X-API-Key"]=clave
    if contenido is not None:
        encabezados["Content-Type"]="application/json"
    datos=None if contenido is None else json.dumps(contenido,ensure_ascii=False).encode("utf-8")
    req=_n13_http.Request(url,data=datos,headers=encabezados,method=metodo)
    with _n13_http.urlopen(req,timeout=12) as r:
        return json.loads(r.read(2_000_000).decode("utf-8"))

st.divider()
st.header("🔗 N13 · Conexión REST entre CONTAVERSE IA y Frame")
st.caption("Consulta HTTP real a un servicio independiente. La integración con agentes conversacionales de Frame requiere configuración adicional.")
_n13_base=_n13_config("CONTAVERSE_API_URL").rstrip("/")
_n13_frame=_n13_config("CONTAVERSE_FRAME_URL")
if _n13_frame.startswith("https://"):
    st.link_button("🏙️ Abrir ciudad virtual Frame",_n13_frame)
else:
    st.caption("URL de la ciudad Frame: pendiente de configurar (CONTAVERSE_FRAME_URL).")
if not _n13_base:
    st.warning("API N13 no configurada. N11, N12 y Failover funcionan localmente como antes.")
    st.code('CONTAVERSE_API_URL = "https://TU-SERVICIO.onrender.com"\nCONTAVERSE_API_KEY = "TU-CLAVE"',language="toml")
elif not _n13_base.startswith("https://"):
    st.error("La API debe utilizar una dirección HTTPS válida.")
else:
    st.write("**Estado de la conexión:** API configurada; falta ejecutar una consulta para comprobar disponibilidad.")
    _n13_c1,_n13_c2=st.columns(2)
    with _n13_c1:
        _n13_probar=st.button("🔌 Comprobar conexión REST",key="n13_probar")
    with _n13_c2:
        _n13_calcular=st.button("📡 Consultar N11 y N12 mediante API",key="n13_calcular")
    if _n13_probar or _n13_calcular:
        try:
            if _n13_probar:
                st.json(_n13_solicitud(_n13_base+"/health"))
                st.success("El servicio REST respondió correctamente por HTTP.")
            else:
                _n13_payload={"datos":_a6_datos,"dificultad":_a6_dificultad,
                    "presion":_a6_presion,"contradicciones":_a6_contradiccion}
                _n13_resp=_n13_solicitud(_n13_base+"/api/simular","POST",_n13_payload)
                st.session_state["n13_ultima_consulta"]=_n13_resp
                st.success("Consulta HTTP POST real completada. Cálculo educativo basado en reglas.")
        except (_n13_error.URLError,TimeoutError,ValueError,OSError) as _n13_exc:
            st.error("No se pudo consultar la API. Revisa URL, clave y despliegue del servicio.")
            st.caption(str(_n13_exc))
    if "n13_ultima_consulta" in st.session_state:
        _n13_r=st.session_state["n13_ultima_consulta"]
        st.markdown("**Respuesta JSON recibida desde la API independiente**")
        st.json(_n13_r)
        st.download_button("⬇️ Evidencia N13: respuesta API (JSON)",
            data=json.dumps(_n13_r,ensure_ascii=False,indent=2),
            file_name="CONTAVERSE_IA_N13_Evidencia_REST.json",mime="application/json",key="n13_descargar")
st.info("Transparencia N13: GET /api/n11 y GET /api/n12 entregan datos sintéticos de demostración. "
    "POST /api/simular procesa los casos enviados en esa petición. "
    "No existe sincronización automática de la sesión de Streamlit con Frame, ni se ha configurado "
    "aquí una herramienta conversacional en Contable IA o Gobernanza IA dentro de Frame.")
